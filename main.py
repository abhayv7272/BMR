"""BMR Trinity Daily Swing Screener - GitHub Actions Edition.
Logic 100% same as Colab notebook. Sends stunning HTML email + Excel attachment.
"""
import os
import time


# ============ CONFIG ============

UNIVERSE = os.environ.get("BMR_UNIVERSE", "top1000")

PORTFOLIO_CAPITAL = 20000
RISK_PER_TRADE_PCT = 1.0

TARGET_1_GAIN_PCT = 0.02  # V6 (backtest): book T1 fast
MAX_HOLDING_DAYS = 25
T2_MULT = 1.06  # V6 (backtest): realistic T2 at +6%
TRACK_DAYS = 25  # purane signals ko itne din track karo
EXIT_ALERT_DAYS = 3  # itne din ke andar exit hua to EXIT alert
BENCHMARK_TICKER = "^NSEI"
ENFORCE_NIFTY_REGIME = False

print(f"[*] Configured Parameters:")
print(f"    - Target Universe    : NSE {UNIVERSE.upper()}")
print(f"    - Portfolio Capital  : ₹{PORTFOLIO_CAPITAL:,.2f}")
print(f"    - Risk per Trade     : {RISK_PER_TRADE_PCT}% (₹{PORTFOLIO_CAPITAL * (RISK_PER_TRADE_PCT / 100):,.2f})")
print(f"    - Target 1 Gain      : +{TARGET_1_GAIN_PCT*100:.1f}% (M4: FULL EXIT at T1 | RIDE: SL to Cost+0.5%, no sale)")
print(f"    - Target 2 Gain      : +{(T2_MULT-1)*100:.0f}% (runner exit)")
print(f"    - Position Tracker   : last {TRACK_DAYS} days signals, exit alert {EXIT_ALERT_DAYS} days")
print(f"    - M4 Exit Plans        : score 0-2 FULL-T1 / score 3-4 FULL-RIDE")
print(f"    - Max Holding Window : {MAX_HOLDING_DAYS} Trading Days (~1 - 1.5 months)")
print(f"    - Enforce Nifty Bull : {ENFORCE_NIFTY_REGIME}")


# ============ UNIVERSE ============
import io
import gzip
import base64
import pandas as pd

# Embedded master dataset of top 1,000 NSE liquid stocks (Self-contained)
EMBEDDED_DATA_B64 = "H4sIAMzPkGoC/6V9WZPiSLLu+/wKPdWZMVOXzXTbOffcx9ACUqKttZBFvUWCMlGlQIwksor+9ffzCAkEWiDnWluCunAPxeJ7eHhEp91Lkat6sTvw/Unx+C5V7f3mWNXl6W9//M8/fc9U8aXgW3lmruLUm6/qLNvz/TrjuRKl5Ue2Tqu//eHanmEz9Q9XAXrGJaCRfaRllb1m6eZvTNNU/HV/1vkhq9HKvCg21d+Yrqv4a34q9ujCcV1nxV5xeZ2WeB3BuGbkOyxU6UmJipyXilXkm2z/VknMoPiZln9jNjO9uYovxdy/ZfsUDezfBl8bOCzwHcdX8aSwQ5HnhRIfX9JqEDpasCT2ndhQ8aiwY13sihrDlMDy/1+yPFX4fiN+Vmhui326r4GdaMxbqCxRoh3Pc0VOZKpofP8+OrXs2VHxp7C3MlO0Y4XBVE3fZryqFbf4oKHRjB13adl29Jsd6cwwI/X8pMTpersv8uItGxkcMywWWjOsA99sxczS697O/RztIgtj2wiTOd7GyzpTjPLYroeV8rzernmZSjDQSQPU0NmlM9t0l63FMhNggL64DWiA3mDG+MtQo0tGr/3gVdPPDDQ33lVN82NJq+zlpajrLkF2m9VNlUnyu6JF89/H7LDDag6TcBIz6s36WHP05zyifvMG82zTAy2rbMP3Gag0Ld9ORNFHes81MUvg+AwKdjiUWdVOnJvW9I4vipvtsVwSfh6aptdgzMs03bevuG048MM4agCDoqwrQbrRIRUTZ66LfYFBKN9BwxL3MpcS/dkMz+ho9eoF8bwdYFyIeeJNn/1M/s+XhnD5C3hmdkxp9TWdBXbMCDOrT1zRsjLnSjvV4ys7C29QwB9bWjQaUIhJyvKLdBG80kF2tBtkJ3tNq/qEbmklGqjGUUPzBjVMaeaqGnJLYtE/1CeARg5z9Wvg6LgX71Lwy+jg9CBwkqhFtPevRQ1mHqTB5bkzy6yi0Y/121ia3veVSyJi80H8vQGJ/HXapUJQ9EnWnNuR40Oopm8ZZgNipKqz9QPrSYhLP2CLBnVZHPg7XlLuMND8kQb+pLGn/z4Oiy0ztogE03qL8U1IlZnue2jndU38hTksueTrY5n2tQ7AZw5kwOsr+vFHJgFo4iGIhCQ4S1Osq+UkaBsMus2PP3kOeYKm6pKvwU1/FwLmH0NveGJe3Ii5H3xf80bOrVPIgGGxsUhcjOD9uKsaGUukfRcrcITCytPdCzj5LrhjL2m+8+yDJpxIM4JQBYEMwS5M9D5/T3eKw1+KktfFZe6vAVegfRv2BKBPucJ2pMhuRWRntTQDhJznsB0ULU/3G5LqNFqDyC7PL0J+Qgk6/gKynsZSvPcoI05/oSHihND8gkHseAnW5T94KykhUEmXg5ce0+6uRmS4e8H7uzK6a/G0/TOOJRG4QEqeoGzo/S9HvFtPSbtU02aQxzwjZLFlq2wvhBuvt5nyLKZ7XEACKw7Zk8CpaaTXwsmbmw7Ze2z/luaKv5/Q9x44Dj3eg+N2EGNFNUofXhKSQtwfD3yHV1b8xPdXM3JZ8ICFcrEOfICNb4w2stgsPwrUxmKzikoAVJ2pH2AEaehJFDdbl4USnao63U28JV6FZosSn8q2Q/eIIYiJhw41WGjJ82PaM6Y6kzA0xTrWtlxvU0xXO0dTsi1cYu5mLILKRwvoUKP7xtWWwGihbxgiYpbdGEkV32ZNX+c5rx4cfWSxIAnB7mhgyw+gdjJOIDFJTgzaLJHlLxxzRfDgVSc95bzt160JDkgm4aBDj1m+WXe13EW8RjbzAmaTXK4yzGPAsyvOuuHEKDZDA1QNesA/u8qfRy54X+flAClFYCXX1kOfEMTwQE0/OTyBoCw2x3U94j8AzZEo+SAA6TLTM0PoWaHPulbbvVmPHXAmM0lt1blQKGaerkEwHQF787YEnamPeY+ektD2vYBGdywxmQfilMY2nVaFSegLlcYIE33db4pG4/Tn0LYdGtP2uAOFfcDaEob1l0HdPMQjSwh4UNgHz7udGJMZS0yOPTNNgzD28Cpmabq5r0QMF94IUNI9mDg6HmC28LIe5StdZ7CrErPBqJg94dnNbMiIn6+wiaIDNPKt+X+BhA8nXcdfgJ12F78zDPAvvrnn+Gqm66j0MfijE6n4w0qAGfZilTsvGRu7Fpkq/kb7psHiefqNHGhVg579ISn4EdIWmCBsvUG8oesBdgbGzIbGgx8nUe45seIVQImWF4QKP04jWI7hzRuEcyDiS5ecp9Fnzhn5MV9bY9KMwvrAIf6RtYbUDQMDbBHaEYl4wL1DHW73vKdA7s66AyGnWxDjaKTkOwhyaCNIbIi7PL/PO2ggmfnhnNYgPyqzonxL76h2kLnuS90DOl8XF3F6ZUjf7blnWMzTwDJoZbOF9J/kGmIujYW+Id76rhSvgMeL+RjwuYcCVvRsCNRlFhOM20K6HLIQqhnyf7jpmLVN13zadNTYygyhgALAnvCDXhaHxhDrUYPpgBiwCmkuDKPP0gHUkSk1qQZthHc1qnS6fzDgNRptrRinPd+dfcVbOXQGk1xd7McgoVdYTPTUIkiCemgI1uU1UEEfp3uqUQv0M0KQold5CjWlF+WBnJyzGpxwW6m7sY0/2Q4UDsvKOs1bYwvvL3a7I0Yr2mvgLfMba+Gt9BcHzCiCvWBPtqpl7yQHZtTtG4F9jz/nof/8jAbgZhX7lxKLWmE0JcekLjFtx7O1O0Sqtg+fF8jF2fzq6HbNDh2m+2FAECJ8cztzI76NFvmzWOJUxWs9aWxomv9N1XK+fle04tc0KHqz0BJ9IeFfjuv3Gx2rOYlpkKrXyFo3sACK+etQnoOtV4BPZgP3lAp6aobdnwbARjELdV+CRzUvR/lFwJIHRqBRTQG3p/RnSr7uSRD1JSQ11oQWQ0MWuxd+gmWMadgqcck3cuVv5n9KbPuhT0YoSKMoiwokHqb79Kd4xyCv+JFuUYhBgye4fYgfZ3YYxXq4UjWiupcUDd8YPwPmRWjPmYHpKbM3vkn7gdDGkwVczDyPhGiJXu73WV/zTYzeZYG7kkJY/2p+FdG2azdxhMZ0SAz8dTXWP+6+TTcjXaWPbtxUn8uwqj5vYqo0he0QiAlvpupqNXTbbHpvm90dio6wvrM4ugvDYear+B4Y/cXUBZl4eI34Aons0s1Q6EcnOwQzI77GI5yRJhQlvqdVtc7cAK61TtZIpVCM/gNMMs5WOvNg1Vm+awJnT/YVbK3dhGADgrYgWIoHUVdGgCwH3kQLZ0WaLmNlNjpQChOu3VUbfREZV20DYfEC8a6wqoJEcfFvbyIQdL8RFtiOw8BJRAV5ziErrlyhzrKP0S0LHDvwSb2jjRwTFBQZWa5TAT19TtTOwX7KPC9e0KV7AXqdhRAsiWcvgVei6SMIeqcke7FLOOwOAwcOs0HTXMJk2nRj0wO+F7nXvkNGr07+Y5F3Rz+hqnUTvjPkl57y+jEOMSmQpKcUFHouynwzSn2mS160DktM9ASy4Ue6HorvAZTCcw7RXroX8YG75qVuRM4Z2kgPkNVYrtPFUbuymYcawDuTcBU4K9HMEbhBfnopeLkZiFzfDi0E86eg3YjvsWbl6WejAoeALeZAZ+pbnpPGLGAHjbvQFnM1Z2aGAn5HtDVLYQ/l2V8U9f0yHjTWLdPzAjOmGd+m+z3P/kPbTbd8W4cYRTNFJoJ1fbNqYDYt32HSIwRezncgHThpO+UCe71ZPdbGTL6500LHoSR6bL3Eu9LBDhwsUnbI++EXPdHwS30iDizu+Ee6YzJ0KadwZLMZIPpxkSK9pfCZ06ihgj/Mhz45m5FlB8Baw9FUom12OIEaB4WDL31LvXi9+AFjMs63RCxAL7bSvSZTaixqrftOwMBYRf5GW3gBz3dFTir0mp0mlDpzI5FUcaRYIkPfOhL9OooyNNk+XODQgIFNzLEuMHpY2fXQppwEJbAaZhQ5gR0ibyXHgNp24R+bNMKyIBJL8yEK7yxmyGZx5BIJlPy1roAjLYsON92TmKFpwCzTofBK2Ak1W6/JtJ6XfEeb0+M60v7OCCf7i69HBQYGFMByBhheCMOZdrz5B20vtbB9b68vohLXBfEJ7XHc7bJ9NZ6soieBTWCHbHOfGFZi814/Zedw0AiRGrq0gfA9zY4GLDMrtGE/ANYFk5QZ5rGfdGM4MxV/V9ax4YM26eMzNrHBtCRUDf5yLK82LqYwHOmvAyvfAaHxpad8P4PFLGBx7KGDFP4IeE1EeauJrlaCcFwW27rEoUldV61Jcs1qY7NumgFbCHVjpCltTTfappLaRmiQ9VnlEId9vWUQ2Qb0d9uEl2Fi67QP51j20oS1ZqT5llbqdMOdhrlcMc/G7x8nypsY4sseA0DM6kzDDGR8V9COsHAaBre4rydPc4CTZ4epfRTDXtqRwzSirOy/qmmz0LC/gQ+N7NdtRP6OMcHmLHyGibAyVaP8CpHJy588x9tk28ObMA6DnI4t0TcgORgXSGZLPRzoWQgJZKwEZJhuNqd7Q7E9/5vMorH3xa/rJJprwpJOWxxiZbFycYl1FXk6l/j7lztRT2PlERWbuqWKWBlR8d0dDdM2AhaClsyv9lfjK8Zegpwe1VGmbfkxdAC+x8jKDPBzMLhBYIKSn007ilQT2uNnmlWV0hdUowLMtGGqha4fq2a2pm0ut8AqPBaNpG0A0BhJdErI6syyXnwd7KsL97MbW3QpuDMUu7z0D++IyZ1o8eD/puTiVPXFert5izO3zT8TmC5m/pZdEsVGgB2ThpCn+xesWpZ2vXsYR/2AvukyDAM+irnju+z+6roUfjVhPFfVCa7BR5oX3e40ysB09SQ0AbcmCXEvM8R03aVJ0LuPlKCxZh9FDstj3U39uhmpZyShML1MjFA6xn3Cvrfk3tz2zBC6WW3Xuxd8PssrKFsTDoCIDhXd4NzdAJNpEz/sPzLgCeGpXN4mUBJSC3VGoFjpmEW0wUArXcOe/UROpxnaYJwyq6btUDPSRXKeWa1FYt7i+FLUg/aIiSF7MFlNoSxG9YQZW34kiKjeFqOGkAmSWDDY1ppqgizeOcXZX9Jx+G+2QSEn1fyVbT69ueCtFoyps8hTzN/wz2h5fRt97g9kZhq0u6fO0s0LTfQnRA9QTXKzsXLATs9+9ig802O1axHQIC4eaFxySp4ryp6un9kmZUWjzXR3Myn3o6Me9LljRtQheK+/IHnGgq+ADezAPENOK5pZ5ACyrOqqOJYTO84z6DIKXAP4I/1NhK3Pucf3tihBObpJsh1kA7BPiHZgxuAM4NVZNRVXT2JH5oHPjj+yE/RlI4AmsnnmDFqAPq415IQ3Ol/GXwx1boIay33xwRX877hLMLd0tI+PWzqYu24wY4nhhCoeleCVHzf5sLCcuyGzQ8rIBWhI+0WH4pxrcJ5f/BR8SWwBcgkPJ+WLSK0isdWxSuZMC0mszfkLaCL/RBR4zhzKJ0FXINf4nlKlqwMfTHyeh5EJsHIDCwmaBY4q+e0vZFNKG/pGjt5iP4eWLYyfOZl9WG8r+00E+mZZPrKaNhxKdZ7uBQN3Y659/3eIRuc2NHAcownKJYSb8aBbM8eUgCbyS95ljzYB8s0nkF9Ftcvq7XuOod/VrXMYBC4LF0BM9ztevt/FcE2D0n3UxtexOtl/3XZ9w57bWMRCMbK3rFb6czbGx/OAOKbYcCxJm9z9ZYIIfGMWrgLLCQjptUxPGAJtLR6qRyfXN0Lzic2hvdFCmf6gQw/FR1o/iKkHLd755+tkrLstkLxumriVo12/WgIHtPfeQOM9B9IRN1tRc39hfgPMO883EDLmrw4/n9PtAJQ4zajfj7kYNNhISugHeg7zKiFdMcdyHvPrkH+XEkIWEJsR3GGbXZP8LUeyyHYJsMp2fY0y7L0DaWkTPUIjZtfJC7fZfnNTxPgoSlMrJqeMOxnqO1yZ8RdKDE22JGsoVi+Rnbo+u6s3nccqPsPCBmy239AGquJh2kegGQnH4w8RnGhycK+WaWLmZ07ih75OubBtC7P8WHSjBT1dYOhkf7XglBQpY/dnA70fp+7NnjfTzy14JCQ2nJJMc3BcN1B+bajc9iQILuMOsgOsmA9x/uN26qPOyyJxpOHBYLxlzlXrq/l1PuiUWbDoZyEDxPzrjbF9SVjrmvUW7CHhBqh4upfcBjUs9Amp4763MRIIsgxtFqn4/Iw5aRkznbbx6Pvzu3eEJUJ+An3SDiUIseMoQD+51WgBV6WPsVwSiwVBFLuUf21xsGFaCeLcVJ+YPrCo4xA+aDm/kxxk6fOObGrjdb0d/Y74ApfZQuJaILDqsE3LtCt1hyLdFvyseD7zfQNIlArwljZpMvc4G5g+LFhf9wmzEDYs8eVDhpNlfotVyt0Rpszj8xfRGmXgZmjb5vyVyD9u079v+QtijjnUw4w2idZFT1DfCg6LOQIYQBTAx8D2/FiPZFtZut4BvpL5lzjLFYOiP7ofBGZ4hXc4tLZurztAkJt1F/j/bLuOWko821l2X53ssUAf7cunFhvY3+FxdXC/Z/v1cKdFYobc2LIwW+vt+ezeqDal3AOR76JS4oEi3K/eDt70cQDL91iEN2LBKn7p/92RAS2h1+7Tk9CDnU2b8e4mBhFVk4/6pXEsJjXUUI9XCcjTVq3TEdSZSS/wE96HreM/IRrF07RsFCBzuwF1KAGq3ChD3sG0mBT4JM1lQ0F5hEdTE8inJbtoIQgTx+439knhbRsaGsHHnUmAXhBkpqn02NDZNMpMI4sXX4NG3rXItmc6ujHTxxPabXvm6CyIVHo454PcVaEELVHuRRYoJC95j57GSdgOMQehdrsD0tDwxQ3tii87pKiy+BzbdLnAxroInmOksX4nmwEQAuoer9oxZje2xxQ0hj2nnYhN+pbu04HNCtoDgiEq5+SRU2QCZe6swOzt+Z7Tunggq1sgRpYThw1itE1z2ui+yLT78kE0Io5WiDZ2MiSPVnZpNZrgI5A8TaLsR5OzbPhcDUgjm81f6y3fv01RliG3Qxq8Zk2vNN1NX9wZa6HPSgKmMawFlufFaUQH277eYpFC+4yWs/3zyH3yRVNejU9BqMfnF4U8y39yOslUS4uZUeoG3Mus2g0n5t4MNZz12vrUWoemcZ6rcyppuzZdzcJA3OvTJNHM/cCLI9HaW3HnTBfkrlhR0izHfT1+BNSeCzg4xiQvtvyhg+pkdESxSNQUDWNKf55lyyAHJ5RdoC0kOD6mpTN09yK0hbmomJu3dGQPtTM7s5WArk7TluZndzYJ/lkoCgJ/zoYP5tleTNtaOnF0TVnL6xprW2VvMDXLdM/vdIlWVmCW5AykCvvIOlR1GeRcLGlXPs/TnbCv5RHJqs7qY31bq8BeEMl8iG0FrNRiX/zMU5rTmwD4WPd8z/ymW8ybqzZtOrbi5GpFLrtDdqAz2pW2D2s+vcn8tIDQpvSQp6+Lr43YnpTaT8ymI0BPtoonRT4qkhJD24yGhc6T5jIVH48fuHpaOGwRWS5es8AI3qvtLnuoe4uAkQcArICfLf9ZUQqHsgMmjtUCio7UDmyP3+ueKw93gRae3I67PsZKT9FzO8/R80PjiJ6NxEm+CXjjmB9/jckYAIjzmisB2i+ygX+VkQ76eSglo9vJKDahgwgyEhvOw6vJbC80dfUJsu8cJ8TD+rTOz/GQHlIgU8qBBBn3zqFkZUT3equt6fOXhTC9nzgEGNpYcFr/clpcPTHXY+JwH9DA7mLNPruu0UwDOrA/sZv6xFaeKQ6NAfXEq5/40UsH/OFrefUUOZjqlCRTlL1vC2GWX3ab+vL1STjcEXtWn4THrUT850jDZwhKy8vPZ0h6iyJbJLvjjDC+7LZPeaFPWfGZ+NRTojkiAvJ0xKzTwX+KgPwsyvdRi5VQbI9irhcke/9Wph/9U/sE25z3PcOKjQvaL+sJO8fCQI+YKxIM5AQ5tDVyObvfQ3i+wD/zt5EDD0+QfbALCbSqN9n49vfTyo+tFQnmp1NRb0+dRKEJc5ewbN3TBRKEoKf3ch+v+7MIQtcGbS2+Bl/Dr+Ko5I25sTBtFX93CHThCZcEX1fhlyHnZRE0BX/wMFruB7/FIipKD4/HpRahhrHg4+5MLUJ0IPQUy2Sx0uhLM1SYZyihOQtt/A+LoUoHBxtpKv4Gf2JPLCTvJlQX/AcvhX9Tdg8U9oTygrauFjw/8JqXx0vW/T3HbsGcFQTfk8SmWhXtCaw78cwF8yJGOSrA21c8gwAqi5yv7xioi1gK2gWnXvF3Pi3pFixMwuUqWjHCOEJ4n6rTPRwKKy44HIYM/5hubsINE6tJYjUC6onOF3cyvUfd7gWU0iJd35/lGdmKRIiv2dDB/aFR2KHjR1roo0NZmRfVO2UklAXVSBhmHoFAReAu8GTLm8MbBgQeeEkHONinR5mHdzVht2gxfM0FBNQvOjjZcbovvO7HrFnkAitMJ4Az8jTurJrtYqTnY9utUUsniNONsHRHa/I48Ux1vsR3gylODDeKADtL+1DCrDNvDpg586sju5MM4ti6Fc1JgeHp4dPuGAt0F7oZPxi4d2IXwO50/2PV4XCf99Ct8IJfymKARh0Ww1xc2uYzgGsyF5dZ+hOeM8TCJWQ+9gaWhAnlzKLnXDgcgympzjexd+jwXzCu/fKNYxonirzY38hFclKRcySKS/TPX/R0XmyZDmxKLBYMC0qphlACYhPf+CLsx4lyFo7p+l4M1QL8HQ0SuuVOwMsxvWhBgR0n3RMn1fcT70ASttqLij6WyUGnTFp6BGOlw9WFHMdfGSQO5qqTF6dNdbWD+Nwxh674u0WLL1h3iu9JDNdsEZrfKeLS1cg9NN+wMIBis+W9MGWTS+AkLvsmxKYDwfRL2s6fTqV0koB4EEZV/9C268a6Sh83jonrQx/Sx3Df3XCm4u+h17tRLDbA6fv2LTGIhoZHD3eTn905usS3FJCHsH4kXOMyK9J9P4aO7lRhgBwtQM4TJSMJz2QuZVhc4aV8d7Hsbzr3hfxT9Szqv1yk/idMd7Qy2MRD88w8FgQJnRChTYsDlWDa3RW3bhjQrO7feE4JjWH6SvkJp/4hiPtzLUpkUFvv2XiqFMB8dJERHHh998lz2qGt+0Ats3XxCPQigpFG8O8VbLTxPoVJbBMYZJYSHf86vmef2LdyqY7UAuhVnb5PKgmXcjFdMPJnSIJ9g3HtxJZA7Gyg38S9rsbz3fB16tFf8KBoJ69Yv1/n5w1RsGmIipOwNg758XwM414FINc0IwtrkqbVthgHikM/8B2bmgdZHYp8OtuT0mhFMQmVUhIGakncWxObisFgTWFOgX+MLa+Pipe9tUemboYeJS5gpWmJN0RHaOfnTCiKT9ABnEaHOf4soqbgGOeKLyITn1lteNYRzdJhy6ssmi5wYeqxHwpP3y2rr4oGu4zSbslsumQTPMRX+jfVPebk6ha7XUG1RC8Bz+szhUN9TmILQlaIv2O9hWy9K3U8FuuyaJgqHsc409N0XaWPiZqaHoEMVZH+m2f6czjI8kshw8vWmYNJjY2uoeBRmpkqPm/0k2cFaBsfXZ/ac3Rpe+ChOzPNr64BFHwMK04vcjw7CiTEROTHiwNduvf0NOrf048SovuPlupBZ50opmaVxw2ecvz1J1d45gxGBuAbz43lIP19JpIxvo50beZc4LuJYTf2l8eWmCP+cT09bGl7IomOfoJBIrLnKCQ0eRrVY99ZSK39RYUT+oYCeUhcWmoUfmuKygXHl7zdisHyxo44suClVZ0/mvzrmfGzqQGn/pm+fLbKAiH74eJf/yvwYXD+638V2dNuubBPDCGhVGQKZnnpURZLnNpo8MxnQf7pzzcQT1S81p/LTvKY+1tD6dnhUOzbJAbqZy85YpTT7SXTkoDa+KDajQfeKpbHsyG8ZEkWg3f8IHuhqXn6yLuTJaVzAa+AubDMKjqz82hVIl8zQ98OnXil+i9pWWSKNMqvDHTfm+vCKPqigCWOZVPv+jNbvGRo08+PHZf3HSZq4vn5pdLj52rW+rRfB17yBTaXouVy0vwOdsBWcLepUuz//T9CWZw3O8d1mg/e85NQNwmvPYQi1VN9GpP8fhDbpgtzxD9Ahe2OTfnmidpY0LuABghNVl/jnul/iuT9kA5uU1KFX4qz2w/sG4HDnaaCm1++5/xByeJHz8wJEjdAp4WVEBx3h0HDLPDmT44aKB6IYvMG2rgEKEciL2Rp6Svte6gGIrP1srwDyxPoTyblTsNJgh5p2x5tWtfQF3yMJikGcxMQ5zjRIac5GGtNXA4Q2BMhkMDTLJ920ecqHh8OIwWeLvfg8CBJp74qI3M2EgI4wPjra3L8I1iz+WkiZh8sQ9otV/GtiBylB8V5wERQTQ34QIXEcywxYFQMxwwARsVw0gMdwqxgHdYDU8Wo8jIBVoqRvp5LdshKo/fca2AvqL6yaOFdKKpuoZsObwYsZt4Tc2xA1nz/A2bmg7G6gLbyg5SXeZv/Oz5wOvwZUUQOCGUF4S0qaTxQlEskmEL3qsKD3dNZEm9+V7IGM/s7scArGTP9IVu+6dnfKI6B+U/32a9uPcxGGwTPGNz2VGVrYuicjyY1BbauM4OmIluv+SbLT/IEyCdc4cAGvh0LAso2JP7TKSayPVPEJgMytTp7YAOMQzUmmEPWfEAlJnh+n9eCdkOwxRgR6pBLK7iZYDCMWES1jwO+H0HpMHIIas1fxjcfAh/m6zOdWQMszEYx6dTb9Tlve6i3M11y+KcSmsRW+jy0jQZ3XmabidIrjQghJJcCXBLJJRE0UQ8qCJk4jxSU/Mc9qRPCmxF1owGNZW7rRo+DR7E9F8BVnUHiyMsqqpvutJRMcRZxhUBAoZb8Mdtg5utQ2MCglOdifXzYsg3CyH3yLS/yCLvaKU/Fdl/dMdDQx6VNO11A+RhMn++zQZgYQqDIxNzz+qUK23xk1VUdrXHFAso8Qui9KGcvaCwlL0hC+EDBURwYfudtIK+Z5D9lesafCSRpvFLE/ymiDEFIfqoiSgtQBks0tKJ/JibsnT+PFI+8nFM4vzoMF4zqrIZKqCz4SzpYay3UHLFNRTu9k3tToSlCuPgaB7FsV8WH4rI5DD47Zoowihpv+/q9dmxGqvgc4AIoPIryhRCMMNsXW04VHcUZ6gl5GC49DJbue1lm75Dw3VjPdeO2Q8kXBEqVWbss7F+xcN/CBK5HiPvbs90XCiMQzX8WUC/FTyqdnG/KdF81m3gD4i6EALMxGxBeWTW8lREuZNmrkO/aus5UB/dSl+IOZ4b6DLgUxM5OvMMeX6b895DF5lyOuE7faNRiEyp/3IFEC1TjxLOpCSpxQlWk22jC6IVbYexJaxpIWHg5Ib1NmL5SBZ5kqC5e736ikK1cH050RCRwEqVzbmqsni0QKl/lzanIVJhSBbe6l5kIiJgFJv1e88NojdbQnJnfAPQ6cFb/kkEYmuIUcZjSMWKyp2qqcVmNN+rYosgGELL2eomrpieMnbCZqRa1P01ofc5C2fobkW1vBYZkgIZVw1xycuzr9hYluophdNH8JMZLCgpju5KAx5gv0WzKjQmPLxlVgcF7MD0DBdIiDQqePu7aLRFaZKEBYJuqDW2au3VAFg8VbQOaOJ1H6J883xE9LT2VPrpzHlGShAENiodGDvz9UiNmsBSWRAGvXFAekQfRYjX3HQx8sVLoQeTrUHI+OCMaI7jIdVxmqfi62Ze66ZIV2oFNVY3VKAj/wy3LKJypUbu/2NnS9eMkMswlfsOU0/8o+D/T8clZaJKOmANd2lWzEWPx0l7MHTXivP7I3geiqTdDYDMW2gB/pVPwHba6U9EqYnOyyVfAfOsEZEakY8Qa8uHZnSuXIuZa1J3dNruzEQ9Il0EqsIDAd5yO492p0er6dP9IRMvFdx/gAtpgVS57IQNx2XuLh4VI4AxEnMr/KGJzcQ9upahnSWq2HNtJB6Y/Iy4PCbl4zS5a/mrja+isO+30UXlSIFbp0NHiu70OAssOaT0Oh21Wpufy6hNFoyKID2bS1JWb8yVS1Hpzynt0nCFUq5iiEqp114kZPWAtR7rFTLqqDLO03vKU7iv7zDk74HumbUj0fZptOsUSH6ggF5kerLMopZsa5sV4HdwoXoHl0mrLOWXHPZTzFVlgCvAPbD54iN0lDCiMNJH+B8xFbFM8jbDf6QZJEU+bqCMIFKyeuE4BKLSAdJ/C5cKii9whyCbhhCBbEoxS+K/dHbaBAc0wB1uRhzMrugbpbe9tJ2DkGkcwGA983FqMdFvtlE24W/oT4tg0KaAZbSmP54F4JnzKZMEa+DDdH9/pxsI3/sDta02FSoodtAUq7+pgawWZZcY0TSfAC1ahdeY36TO3NcZCeklGw6nuHseNbDoJEJ3hB0HmHoupMlqUve0plp++yQjVFQW1OmVhi1PP0Xt2Oex80yCddKPtH7yWjrrJ/Gu9oG+aOhGUu63zE/mahakvXrbXrqKofCSuSpMX4HZ00qDDAHHOdKr6GUFqK5rzDDcf7rG4HPJTngM1FNP9DM+iqZo/FkePYNZZwrWMYNptlc7pvXFKoEsT5IkNUXzqAVIllCY5QqA0O0tiIPIewU+YZBoRlKh2cbcIN6luVe7e0mVWRUm6/poBeyL/WVavAlYpjmxQJ5+zvCoGzsf13dYoltUtJHo9FM3tm8yimlcEMtmk59yXqHtdTFeu0Okcz/4G+FMp4ptBWoq1JejzlPcILREJuhGlVA7zVWKwMBJJj9FxAyGyfex6O+CZJjRyKPBk3PsSU+zCuXbT+i4jUr5qfIAzAlCZSjewNu2db0I9OxRXlNKdosRrwpwX9LZWU28k13jxUuDES6XZCX40MhbRkfpQSlQ6Vf+QSE0o1XpBCHW6fh/at4wSih+K7YfoeCjBaXfCjVHgyLoNLfg5O6y/IEFor5gApGhDp9zhwMyEKwb68aDTjuWJrkSotvtsmI6+081v0fGvvJi03Z/pNqkQKv2nqEXRBsY6HXy253PYGD+zt/Ec2mjlzSm2Gp324vD1gBXcHcdKUMVJMNg86ufoXnWR5JxKH1dLEms+ZZHhC+jvo7m9hvDco1Wkxsb9CnhxBMpZqvjqJxwMUE68jKgESqjioSnfcMUMd5QFsCI9EshY/UNOkViKG41VHIxdTY35Lsv3HBZzWq7pjr78ztGvmHkOA9qeUptzXpOMmg5DxSxm5BHFpL7uuUMCmOSJhB6LJAsw323BBjbC++I4prmR4FjWvOZ07vmh9Hv5Ng+yrtPAp2qdUQum8y2yZQNm/qvK7r4SpGNGscTo3EbwyHZJ7OpLiditBXlDnG5wDRPwqkr3b+KE4jZb5w8GC6irMp4lGpNMcU5faio1A0aetBQwE5lXBClUrQB8ONYZx67ToqDj59s4OvnD/xglDLwOK0ubQ9eRlbF3Ad7zTVX+08V3+3JTGrl7hqVrTFBcF8hv/I7Qj80nFtFGapz+oJB6WncOEQ4Nw6PILX7Z7+ksprw7gmWf8U1j8xtF6CkIC+VaiNoDclyDJcViOm8GRxqeZkzZinClT50EidNo7XEvgTzYpgpdU3yvsphnM90R0F7685L/9KBdSTf++OSAUQMh39HEPFCqg4IyLvtGWCWl/g8O3ncZ5bdDa2yLncg7Kt7vlQyIrVXo6xTajbcnqPDhlLCOcoshM7Kcv8u0+09sUcc2/DpBqmhg176lmw9yPRzbXZgegb6nE7WK4ORHbhLZUKFUZtI9Vtn6UWsqtmOqOg/epspgvNxK+ppSoHSvmADf9xnqmqagN9u73eOiLNP2WOxEPU+JIkVXi3IbiYehQPwAYqtIGNBGDZ0sCcHzVJmynAxxxCHzImKopoWShutkb9t6hNaBsDTl6eFmt0dmE99JAo9Dctvi8sphuwWxDQmUbc5g572W/59K+sCF465RncAWPT6WL9k+HV5RGcih/acro+jaf7sjpxLdF+4tvkd3fxPPF2nsKh4Umcb+UNOBoyZBr6ZwEosyVfj6fGWq5OnJXoISZpqa/PiRfdCVPo8ftk8cUAUkmO6rSQ49JhTVA9554tm+J2fpct3QlFudaA6B1nDGtBIi/CEpg7fERhSELWZ0yOCOP4AXintnVVnrbJi3Lz2LLFm5KIH/Ki4TyAbPgS/nCW0oLX+bH6kc2f0KW0vR6vI3cT9BCFMkGz3KvrQD0q7Lr/bX4OsjTf8emjHp0uXvd1o2NLltotLt2qFCVcBj88o5eWYam6tLaQ8pz/yFv91uXy6ZrVlsOccaLnn2suUfbYrXWP9igiw3W7paqBUEN2JhyUIoKwKDpvp0SH8p+lLC8daoKiC45ZEjRt6K4bXqMt2I0gmw3abOMS5Nw2lgRyI/S/uJwR9eZj8oCd/I+Nu+qMQhY7rFre8UL+0FbIbId9CH7J0cfhnzGyI3OCDLDCSZw316k4Q5Tj4rFlkAP1EhEBEGIgWy7tcVM0z0ttjwV7qO1N6ko/kQS9+JmRsAGM7M7qBcKSmRF3HdWwKPJPQo1T4zmCYmlVdRnzmnEPRUWpOEpvt2G+BLKO1euuMzW5gzOwbie/qaUbkLKvg+dGb10jeT4rnqc0qHRxtlcZuGfp7vZ1Mzm4r1z+lLVeTnckvC2hjskumI0MVzmleH41X04haONGkLNpgbcTb2ARwFiefYyzOCk32cWejMas90QIL2I5/pfATtRzZVQ3o7PJchWnSE3vcx+9uszA8FBnmtRQfmkCq6qs/ZoT2CPeLjPPv6oom2PRfrd9qfqXukumIx3aBsqSte0/3J23NhEUowKflxx+U1Qtf2S7cFMxIaapVW00rw+0xfSsPh+0xpLq0ggMZTJTEibp9sbcnHTYnvbD53TPU7f3vLqb5teuDZRvHX4uK9R6ID38mt/Z6mN6bvwLnle/bxd7pIRUdb+7vsA9CoObwL8AryqcG46yl/XxlJJCjt+4lKqE3ehyKARWVvCQwSzvfnI7hTt+Hojhl+U1M4oeWvm3k8L+rv/3RtPfS9SP39nxDX61Iwf7YjQ0Je8kKa7fd/RbFueu7cVX//V1Ur7bWZgxf83SL/ocl7qG1f/UNT5E3UWaEYvwZAA3HTwB+B4hCdXO6NvAX874DZEVP/O+BZ1Ylo3YL9zx+uT2P7nz+UXUEjq7trmt+Cs9+/y4x5PAyV5+7BMxknwfcNtfQhF6R6mKxBZf465G34aHzeSLCvjDCZR8CjiwWM8vgm7ma79akGMH2PkIrrPL0BuKVNcB/kboiL0PogEQQMYESBQFiWe6r60QfTKPcAEpm98P1NwYA+rO7HGJOiUT376pCdXd0+pLuA64yvS+06ufHWg9R104GJQRcs5ifeibeONU0XncbmHBgytP1GSbCbbqC2jzJ3upLsZut5Xgy+Bij7DR3APYc8ejCGQVYt22yy6dUyZrIuCB7awwZ9GLyvXG+zDcWeDVgr5RCQH+J1+AmyZDNM2UboU2r9DLO/KQthHLw2t3H2YSO8VG42its2rk759qFhJoNONh88q1uDQJhN2di8A4Oy/JptKnFOIJd1tIYg6Ya3UAJn/Yof50jBAC7o14IdBVQSaA3qzbbHAJq90BkhZe/jgojRrXKsrKmK+9W1nCLd7px/3kcLfXGzl0p1y1/ppqNp1gL8zDG/XcAnCcqkQyLJBdhLj8NQJJaxAFQNT99Sl5/42wZWt7KAfVEqf2/ujhia1hncL/YuKoDv5WXW40tAtQGkABk6L94Dn4uEALW5NrF7mdxVfL6PB5k3byqGHeiC1QEI6vVb1t4s1AcQFA/9/Facfbw+UCjv9wHBNjf7jA99jkm2/EgVh2DofNOo4LKESKEtwG7xgD6YwwyQ5Tbnm6sb3QYATbryj1JUzxWA/07Xkwx10w4ZFEtW8v3gj7D1XfoZhi2ctr5+fWLfRDUa9oP/uqNan1xKZWM/dpTJ1uzffmnWY2xqnqDssbY/oOqzy+3BPTCpioUmFu39NsGAC7pqHrOzGKjs3QfGyN7foN6z3dCv6NrCHu38glKsgA9ftpxQtA4VGUGHcioyUg8BaPAGDbYEY+QvsHMVg3/Qhb19QNPVbHGogeXp7iUbYgRnFpqGBdbPX0uIdystRZNjY3Dm/p8JIzMgfyv+fRTRg/bsZR9YJDOzXOQy0yWTnCpIDwE6Ogvn/oXjHFhZTSrcADBd/USwG+giGa0ZgqLQPEEpIjB/ya3rg1JK8HeAUqb8X5dLdi+Jb32UAJRLOWzXh9MH4GzPpFpqAKZYLTzRxtvsgyayZADLj6+iXsBmiOldIdt2/C1rTtnLc24DcA6B5YM/hU9+bIlmyh9Fvc0EMVIcpz3v18ehLHISx0ACHXFxk+0LHFIa9heF0X3bSjalidDCQvcF+jtvrzkafZ20l1xjWru5T8KZYLsf97wJ5nqBQzJwR3f9kYUtyaE5FNIHDxNx/BPUuyuP4vzn/vrO3h6KZ1ghC1igyidFFAxWHNuFijBu4KKEplLuzQnTsc09HGg1NJ9XCZmYZfpTWR3zVJx1aCK5fXh7IVLo2D57n548z/VJzeyKfJTZPY8FSeiB1Pd7fjiWVFL2JyezYEgswVd2NOppRafFqAiPuJtjADCmbPaEOllTRTC4vFGaX6raD2HIW750QuHlmtyiIbDECiyQ6P64PWdS9YACoV1/kO1I90KJEhNtoH8A2Af3CmI8rIv6nrkVEKcfADUri7/S/agNH/gOZoCut2SHAjKRbqpKRwAj24O92wJGYLICwuGcIz+A4tOBPMAXJdUlSfnVWXIa7lWewxgl/bmiXCT27xMZIl766zj0slC3bCNwVh2XZHKCCN6MJPQIAFYHH23dwEEYw8TElZs0vfMyYVgHash/tEl50/CUgcbK1wl3NZS9J3xhxw1BWEI9ltlWViWmjDbp93w5H+wfQ6MMnhaztUaosFc6aLlInCi54ERHKq/XJrYP44j+oxdkZk35wuGCGTTN72TRd2v39SFd5smpo12Ohwz70Kfr/ejWPqgOhW74uz3b3UeJIo04lwg45EJriTR1cfi/puJfWlm8j/QwNl07otTz1kM7555PdTL2bBobISleJkY3ALUU1EhXR18ZA+OeQLgUedOEREXkROZ0myndB14x2QW6DvbB2Y1YG8aos4+UqnXvDkIvitBsVcEWzLMxTF2DBK2qAi+hcAnL18W2EB72zX5lB8uyGWmJapuJekyjVA5Al0UrT4AOSufI8hfMNclb2RbvdEagXotcij4k3knZi9Kt6V3zMAgPQeqFN44Q3aG4/ccIQmxTfUSJIKm0GJUNAv7coTilCkjnKxf70LAM6Vg/wLE2dKpfXmw4Ko+jiIm8cawNPC+zcbUH4CghWvZBppiPt0hxcFbBeBfx6cvRrAHIkH03adnw1u/pPoX6bRTs2GzEzKENV5WRDBPe3bCdHpM/CQKlmx4/Bl8edzzxOC1hwQ2HRGO61peBDJuHAdsLv0T6SnewSFD8laKfRALe3y1eUiWxISqIE0cU/2f1MW/umejBJGGCQRzL405c3TfikVDtXFV8KnPTVWgXzG8Lh/d6mkS+XO1jVew6IaGhdqmAyJwgacN+KDtuCCn22TdB2xc5wX7lgwsAWGEBATKXVcwmxPSSGZY0b9kH32ylZSuy20Zsuyb0JyN/d2LsS2E7CmA61z4A4MOsDam5ApYtxHurRb88tpPBlqEcK77bo/29tVnGHsX6lvG5HtYlRbXX4LOM5sDbOSnPVD1TscDn4nSEPmwef6PbZtmvrD67SQMwYhp+1Wk+vRorTMUpvWTT9wFciGPhoK5cJTqNhCS+M+mQ/cU37ZbwuRzXLbDG5jPbcSNV+8q+zpurrsVxwXRQUmjMpkvvVY12BNFRSisdgnpiTxRN0PgP3rkIedgZE9CWSACU8FOBNAlM42tgpR9L5qIk3mEMqs88EzeXvRJtHV85GV2Dek/Aw/uU7cu83KllAzyWF9BpPhE6FpfkwPaQh9jirDweYGgqGs+BKIPoFHbfXOXJDjYC+nZopoCYidTeV3k/+RDwN1uVX4rcudbNJG4qTt4yCuAWsDctmAZouy3qIO/PGfH8geI67BlmNqGIuyD5T0zSlM8LnGYTXetU5elDeSzUTAbfWIMkoPJRWso31QikDRlGgHs6Hc12u5YYhuEF6Nvl+rIhmAX1b09XIWSFUxSDMDC/nm1BtMKJfs7K9A6heNEz1cMihJ9Ue2WMhTUWzVT6GFPYGpSjrflYXV7n2UuRDYDA/ta24sZkLfuRp0O8pMUQd9oXDZSXHfJU2DWUNsDprPDQCHTMXUnJmlT1bl4Wx8MQkCPYc32HbSAfdDOIMR2ltKUg8w6DrKNDgNPH+d632yPBPQwTbmdEe/SgmnJTDYpuzTSoTiwWOt3sjrSdOtld01ywFWDTd3TgAamAtzOd7kvVUqoBpdOdqU3qltj/kulvfTRPlEzTUtDypUT+ENwcjB25BPkmVKY0NccZj7a8QBBEq9fZrFTVaje83pR4R3fCa3Rz1vhGn2bGNE7oQLkn3geYNSc3tO6xnyE4yCcqImYS5DmVbwDQYnNxKAYjessgN7fXVWQmlwa4K4iBUCCf+DmNfJDT4LUbVHuBpm7DRSynyiiRf9RbAEpIJYTZmQH14l3K9aEQrwQXiy6hQ5DKach7lZCR2QJGKR9/vxRNElCIprAYjFpp9lxzfF3VsreXvFhfX7I8AC0HVdNBujt7sRpl29MkZDtKQCQJXg4HfDW6VcqR65kVr3QlZWct7+ZXaHboUKUQc6UyqhRNUrvMyezZp6cRcC9RBZB3HAEI6HYz2c7lpPK0xas5NFrTdWSAWyRFD4JprfbF49DvZE+kWyI2R1xIPrExh6acG+17da6uD0/bbSsWwOJTnK/iVrUTP8jAzGBsSXOI4mAymL+Nm+Kak5i6ZaPV/JhSjanDKEMRpOMnpNoFbF4cN+IMNzliEzEuIEZBSFU3gVcdxHRM7r5r7jPlENCOGh47d8z1AD0ScpL8PCHm5DHSHpxvUTUtNFdsyzvqwneNlSk6SxdBnxTjlMqTA+7rlIHnexTH1QqxoTz9gpCiOFpRUhDnOeUfI1C+GH9RQmrlwwBUwKEF6Gbh9mEj3frNohqk4lGhO8XhGbszP4zH/C+NtA6dHei1di7R2AipgNNhjaZGSQ86ZJYrk8FgLmx3/HCsxRLcMQRCey6v1CWzhUIwQqFO0QJhxFrY2Dmd6536kN5ZwIcFfK3zQeoeZGRFTksIIuqrWKcXKtdId2ENNU3g0dCk0Zk+rdhsM3kJjiiRvxtx1RLbMSi1Ndo2pf7ENs9hKBCuJaEXJBj0sdwfjuXoFouWxLEZzhwYQscaa/Wan5Q5RAfFBynIyw4HWb9raFBLMjbRj3eKHuWQzKMpWOSh+LNY1WGqvBblPusUeRjlHZ25MDxndOhG57s8E+HXbg2jPoIXM1JoOqXKw0NurbIRyaWzgOkUBdf5ga+z+r9SSX+HUadPJCtgLDPtnKs6dMqljxaKS+ogPeiRAujDehO/hgwrDL+65GUx3vFwFdE45XefSXUdS0OKH/bdRgSSr6wA8QtU/nFAJuq6ZRHucV/DJtLz48s54zkTKSKWaPM00Pv/B3ywYk/4uwAA"

def get_universe_df(universe_type="top1000"):
    raw_csv = gzip.decompress(base64.b64decode(EMBEDDED_DATA_B64)).decode('utf-8')
    df = pd.read_csv(io.StringIO(raw_csv))

    if 'symbol' not in df.columns and 'Symbol' in df.columns:
        df['symbol'] = df['Symbol']
    if 'name' not in df.columns and 'Company Name' in df.columns:
        df['name'] = df['Company Name']
    if 'industry' not in df.columns and 'Industry' in df.columns:
        df['industry'] = df['Industry']

    df['Yahoo_Ticker'] = df['symbol'].astype(str) + ".NS"

    if universe_type == "top100":
        return df.iloc[:100].copy().reset_index(drop=True)
    elif universe_type == "nifty500":
        return df.iloc[:500].copy().reset_index(drop=True)
    else:
        return df.iloc[:1000].copy().reset_index(drop=True)

universe_df = get_universe_df(UNIVERSE)
print(f"✅ Loaded {len(universe_df)} stocks from master catalog for scanning.")
print(universe_df.head(5).to_string())


# ============ INDICATORS ============
import numpy as np

def calculate_technical_indicators(df):
    data = df.copy()
    close = data['Close']
    high = data['High']
    low = data['Low']

    # 1. Bollinger Bands (20 SMA, 2.0 StdDev) - ddof=0 for Pine Script v5 mathematical parity
    bb_mid = close.rolling(window=20).mean()
    bb_std = close.rolling(window=20).std(ddof=0)
    data['BB_Mid'] = bb_mid
    data['BB_Upper'] = bb_mid + (2.0 * bb_std)
    data['BB_Lower'] = bb_mid - (2.0 * bb_std)
    data['BB_Lower_Slope'] = (data['BB_Lower'] - data['BB_Lower'].shift(3)) / (data['BB_Lower'].shift(3) + 1e-9)

    # 2. RSI (14-period Wilder smoothing)
    delta = close.diff()
    gain = delta.clip(lower=0)
    loss = -delta.clip(upper=0)
    avg_gain = gain.ewm(alpha=1/14, min_periods=14, adjust=False).mean()
    avg_loss = loss.ewm(alpha=1/14, min_periods=14, adjust=False).mean()
    rs = avg_gain / (avg_loss + 1e-9)
    data['RSI'] = 100 - (100 / (1 + rs))

    # 3. MACD (12, 26, 9 EMA)
    ema12 = close.ewm(span=12, adjust=False).mean()
    ema26 = close.ewm(span=26, adjust=False).mean()
    macd_line = ema12 - ema26
    signal_line = macd_line.ewm(span=9, adjust=False).mean()
    data['MACD_Line'] = macd_line
    data['MACD_Signal'] = signal_line
    data['MACD_Hist'] = macd_line - signal_line

    # 4. Exponential Moving Averages
    data['EMA_50'] = close.ewm(span=50, adjust=False).mean()
    data['EMA_200'] = close.ewm(span=200, adjust=False).mean()

    # 5. Average True Range (ATR 14)
    tr1 = high - low
    tr2 = (high - close.shift()).abs()
    tr3 = (low - close.shift()).abs()
    tr = pd.concat([tr1, tr2, tr3], axis=1).max(axis=1)
    data['ATR'] = tr.rolling(window=14).mean()

    # 6. Relative Strength (30-day return)
    data['Return_30D'] = (close - close.shift(30)) / (close.shift(30) + 1e-9)
    data['Vol_SMA20'] = data['Volume'].rolling(window=20).mean()

    return data

def calculate_benchmark_regime(nifty_df):
    c = nifty_df['Close']
    ema50 = c.ewm(span=50, adjust=False).mean().iloc[-1]
    ema200 = c.ewm(span=200, adjust=False).mean().iloc[-1]
    cmp = c.iloc[-1]
    ret30 = (c.iloc[-1] - c.iloc[-31]) / c.iloc[-31] if len(c) > 31 else 0.0
    return {
        'cmp': float(cmp),
        'ema50': float(ema50),
        'ema200': float(ema200),
        'is_bull_market': bool(cmp > ema200 and ema50 > ema200),
        'nifty_30d_return': float(ret30)
    }

print("✅ Indicators Engine compiled with ddof=0 standard deviation parity!")


# ============ SCREENER ENGINE ============

def yf_download_retry(*args, **kwargs):
    """Retry wrapper for Yahoo Finance downloads (network hiccups). Same data, more reliable."""
    last_err = None
    for attempt in range(3):
        try:
            return yf.download(*args, **kwargs)
        except Exception as e:
            last_err = e
            print(f"[!] Download attempt {attempt+1}/3 failed: {e}. Retrying in 10s...")
            time.sleep(10)
    raise last_err

import yfinance as yf
import math


def simulate_exits(o, h, l, c, entry_idx, last_idx, entry, sl, t1, t2, plan='PARTIAL'):
    """Forward-simulate exits. FULL_T1 = all at T1. FULL_RIDE = trail at T1, all at T2. PARTIAL = 50/50. Returns (legs, t1_touched)."""
    legs = []
    rem = 1.0
    cur_sl = sl
    t1_done = False
    t1_touched = False
    end_cap = min(last_idx, entry_idx + MAX_HOLDING_DAYS - 1)
    for j in range(entry_idx, end_cap + 1):
        oj, hj, lj = float(o[j]), float(h[j]), float(l[j])
        if oj <= cur_sl:
            legs.append((j, oj, rem, 'SL-GAP')); rem = 0.0; break
        if lj <= cur_sl:
            legs.append((j, cur_sl, rem, 'SL')); rem = 0.0; break
        if (not t1_done) and hj >= t1:
            t1_touched = True
            if plan == 'FULL_T1':
                legs.append((j, t1, rem, 'T1-FULL')); rem = 0.0; break
            if plan == 'FULL_RIDE':
                if hj >= t2:
                    legs.append((j, t2, rem, 'T2')); rem = 0.0; break
                cur_sl = round(entry * 1.005, 2); t1_done = True
                continue
            legs.append((j, t1, 0.5, 'T1')); rem = 0.5
            if hj >= t2:
                legs.append((j, t2, rem, 'T2')); rem = 0.0; break
            cur_sl = round(entry * 1.005, 2); t1_done = True
            continue
        if hj >= t2:
            legs.append((j, t2, rem, 'T2')); rem = 0.0; break
    if rem > 0:
        if end_cap < last_idx:
            legs.append((end_cap, float(c[end_cap]), rem, 'TIME'))
        else:
            legs.append((last_idx, None, rem, 'OPEN'))
    return legs, t1_touched


def tracker_status(legs, entry, cmp_now, entry_idx, dates, plan='PARTIAL', t1_touched=False):
    closed = [x for x in legs if x[3] != 'OPEN']
    open_frac = sum(x[2] for x in legs if x[3] == 'OPEN')
    realised = sum(((x[1] - entry) / entry) * x[2] * 100.0 for x in closed)
    reasons = [x[3] for x in closed]
    exit_px = round(sum(x[1] * x[2] for x in closed), 2) if closed else None
    is_open = open_frac > 0
    live = realised + (open_frac * (cmp_now - entry) / entry * 100.0 if is_open else 0.0)
    today = dates[-1].date()
    entry_date = dates[entry_idx].date()
    if is_open:
        exit_date, exit_reason, fresh = '', 'OPEN', False
        days_held = (today - entry_date).days
        if plan == 'FULL_RIDE' and t1_touched:
            status, action = "OPEN (Trailed)", "Kuch mat becho - SL cost par hai, T2 tak RIDE karo"
        elif plan == 'FULL_RIDE':
            status, action = "OPEN", "Hold karo - bechna kuch nahi, T1 trail / T2 ka wait karo"
        elif live <= -1.0:
            status, action = "OPEN", "Hold karo - T1 par POORA exit karna hai, SL paas hai nazar rakho"
        else:
            status, action = "OPEN", "Hold karo - T1 par POORA exit karna hai"
    else:
        eb = closed[-1][0]
        exit_date = dates[eb].strftime('%d-%b-%Y')
        exit_reason = '+'.join(reasons)
        fresh = (today - dates[eb].date()).days <= EXIT_ALERT_DAYS
        days_held = (dates[eb].date() - entry_date).days
        if 'T2' in reasons:
            status, action = "T2 DONE", "EXIT karo - Target 2 hit! Poora profit book karo"
        elif 'T1-FULL' in reasons:
            status, action = "T1 FULL EXIT", "EXIT - T1 par poora book. Plan ke hisaab se closed"
        elif reasons == ['TIME'] or reasons == ['T1', 'TIME']:
            status, action = "TIME UP", "25 sessions pure - CMP par exit karo"
        elif 'T1' in reasons:
            status, action = "T1+SL", "T1 profit bank, runner SL par nikla - closed"
        else:
            status, action = "SL HIT", "EXIT karo - Stop Loss hit. Discipline se niklo"
    eff_sl = round(entry * 1.005, 2) if (plan == 'FULL_RIDE' and t1_touched and is_open) else None
    return {'live_pnl_pct': round(live, 2), 'status': status, 'action': action,
            'is_open': is_open, 'fresh_exit': fresh, 'exit_date': exit_date,
            'exit_reason': exit_reason, 'days_held': days_held, 'exit_price': exit_px,
            'eff_sl': eff_sl, 't1_touched': t1_touched}


def run_screener(universe_df, capital=500000, risk_pct=1.0):
    print(f"[*] Fetching benchmark data ({BENCHMARK_TICKER})...")
    nifty = yf_download_retry(BENCHMARK_TICKER, period="1y", progress=False)
    if isinstance(nifty.columns, pd.MultiIndex):
        nifty.columns = nifty.columns.get_level_values(0)
    regime = calculate_benchmark_regime(nifty)

    regime_desc = "BULL MARKET (Safe to Swing Trade)" if regime['is_bull_market'] else "CORRECTION (Caution Recommended)"
    _nc = nifty['Close']
    nifty_mom_map = {}
    for _k in range(30, len(_nc)):
        _dd = nifty.index[_k]
        _dd = _dd.date() if hasattr(_dd, "date") else _dd
        nifty_mom_map[_dd] = float((_nc.iloc[_k] - _nc.iloc[_k-30]) / _nc.iloc[_k-30])
    def _nifty_mom_on(_d):
        if _d in nifty_mom_map:
            return nifty_mom_map[_d]
        _past = [x for x in nifty_mom_map.keys() if x <= _d]
        return nifty_mom_map[max(_past)] if _past else regime['nifty_30d_return']
    print(f"[+] NIFTY 50 CMP: ₹{regime['cmp']:,.2f} | 200 EMA: ₹{regime['ema200']:,.2f}")
    print(f"[+] Market Regime: {regime_desc}\n")

    if ENFORCE_NIFTY_REGIME and not regime['is_bull_market']:
        print("[!] Nifty 50 is in Correction/Bear regime and ENFORCE_NIFTY_REGIME is True. Screener will prioritize cash preservation.")

    tickers = universe_df['Yahoo_Ticker'].tolist()
    total_tickers = len(tickers)
    chunk_size = 100
    chunks = [tickers[i:i + chunk_size] for i in range(0, total_tickers, chunk_size)]

    meta_dict = {
        row['Yahoo_Ticker']: {
            'symbol': row['symbol'],
            'name': row.get('name', row['symbol']),
            'industry': row.get('industry', 'Equities')
        }
        for _, row in universe_df.iterrows()
    }

    actionable_signals = []
    watchlist_candidates = []
    recent_signals = []

    for c_i, chunk in enumerate(chunks, 1):
        print(f"[*] Scanning Batch {c_i}/{len(chunks)} ({len(chunk)} stocks)...")
        try:
            data = yf_download_retry(chunk, period="1y", group_by='ticker', progress=False)
        except Exception as e:
            print(f"Error in batch {c_i}: {e}")
            continue

        for ticker in chunk:
            try:
                df = None
                if len(chunk) == 1:
                    df = data.copy()
                    if isinstance(df.columns, pd.MultiIndex):
                        df.columns = df.columns.get_level_values(0)
                else:
                    if ticker in data:
                        df = data[ticker].copy()

                if df is None or df.empty or len(df.dropna()) < 120:
                    continue
                df = df.dropna()
                if isinstance(df.columns, pd.MultiIndex):
                    df.columns = df.columns.get_level_values(0)

                df = calculate_technical_indicators(df)
                meta = meta_dict.get(ticker, {'symbol': ticker.replace('.NS', ''), 'name': ticker, 'industry': 'Equities'})

                i = len(df) - 1
                c = df['Close'].values
                h = df['High'].values
                l = df['Low'].values
                o = df['Open'].values
                vol = df['Volume'].values
                vol_sma = df['Vol_SMA20'].values
                ema50 = df['EMA_50'].values
                ema200 = df['EMA_200'].values
                bb_mid = df['BB_Mid'].values
                bb_up = df['BB_Upper'].values
                bb_low = df['BB_Lower'].values
                bb_low_slope = df['BB_Lower_Slope'].values
                rsi = df['RSI'].values
                hist = df['MACD_Hist'].values
                atr = df['ATR'].values
                ret30 = df['Return_30D'].values
                dates = df.index

                if np.isnan(atr[i]) or atr[i] <= 0 or np.isnan(ema200[i]):
                    continue

                cmp = c[i]
                is_uptrend = (cmp > ema200[i]) and (ema50[i] > ema200[i])
                _pend = None

                # Check regime enforcement
                if ENFORCE_NIFTY_REGIME and not regime['is_bull_market']:
                    pass # Enforce cash mode
                else:
                    # --- 1. ACTIONABLE TRADE EVALUATION (7 CONFLUENCE RULES) ---
                    if is_uptrend:
                        # Rule 2: 50 EMA Slope Guard
                        slope_ok = not (ema50[i] < ema50[i-5] * 0.998)
                        # Rule 3: Relative Strength vs Nifty 50
                        rs_ok = True if np.isnan(ret30[i]) else (ret30[i] >= regime['nifty_30d_return'] - 0.025)
                        # Rule 4: Lower Band Touch
                        pierced_lower = any(l[i-k] <= bb_low[i-k] * 1.005 for k in range(2))
                        # Rule 5: Lower Band Slope Filter
                        bb_slope_ok = not (not np.isnan(bb_low_slope[i]) and bb_low_slope[i] < -0.012)
                        # Rule 6: RSI Reversal Hook
                        rsi_oversold = any(rsi[i-k] <= 40 for k in range(3))
                        rsi_hook = (rsi[i] > rsi[i-1]) and (rsi[i] >= 28)
                        # Rule 7: MACD Selling Deceleration
                        macd_turning = hist[i] > hist[i-1]
                        # Candlestick Confirmation
                        candle_ok = (cmp > bb_low[i]) and (cmp > o[i])
                        lower_wick = min(o[i], cmp) - l[i]
                        candle_range = h[i] - l[i] + 1e-6
                        wick_ok = (lower_wick / candle_range >= 0.20) or (cmp >= (h[i] + l[i]) * 0.50) or (cmp > c[i-1])

                        if slope_ok and rs_ok and pierced_lower and bb_slope_ok and rsi_oversold and rsi_hook and macd_turning and (candle_ok and wick_ok):
                            entry_price = cmp
                            recent_swing_low = min(l[i-2], l[i-1], l[i])
                            stop_loss = round(recent_swing_low - (1.3 * atr[i]), 2)
                            risk_per_share = round(entry_price - stop_loss, 2)

                            if risk_per_share > 0:
                                trade_risk_pct = (risk_per_share / entry_price) * 100.0

                                # Noise Buffer & Loose Stop Filter (1.2% <= risk <= 8.0%)
                                if 1.2 <= trade_risk_pct <= 8.0:
                                    target_1 = round(entry_price * (1.0 + TARGET_1_GAIN_PCT), 2)
                                    target_2 = round(max(bb_up[i], entry_price * T2_MULT), 2)

                                    t1_gain = round(target_1 - entry_price, 2)
                                    t2_gain = round(target_2 - entry_price, 2)
                                    avg_gain = (0.5 * t1_gain) + (0.5 * t2_gain)
                                    rr_ratio = round(avg_gain / risk_per_share, 2)

                                    # FIXED: 1% per-trade risk (param) use karo, stop-distance % nahi
                                    risk_capital = capital * (risk_pct / 100.0)
                                    suggested_shares = max(int(math.floor(risk_capital / risk_per_share)), 1)
                                    suggested_shares = min(suggested_shares, max(int(capital // entry_price), 1))
                                    total_investment = round(suggested_shares * entry_price, 2)

                                    score = 0
                                    if rsi[i] < 35: score += 1
                                    if vol[i] > vol_sma[i]: score += 1
                                    if ret30[i] > regime['nifty_30d_return'] + 0.03: score += 1
                                    if rr_ratio >= 2.0: score += 1
                                    if bb_low_slope[i] > 0: score += 1

                                    rank = "A+" if score >= 4 else ("A" if score >= 2 else "A-")
                                    g_score = 0
                                    if (not np.isnan(ret30[i])) and ret30[i] > regime['nifty_30d_return'] + 0.03: g_score += 1
                                    if rsi[i] < 55: g_score += 1
                                    if hist[i] > 0: g_score += 1
                                    if vol[i] > vol_sma[i]: g_score += 1
                                    guidance = "RIDE" if g_score >= 3 else "BOOK-FAST"
                                    plan = "FULL_RIDE" if g_score >= 3 else "FULL_T1"
                                    _pend = {
                                        'rank': rank,
                                        'symbol': meta['symbol'],
                                        'name': meta['name'],
                                        'industry': meta['industry'],
                                        'entry_price': entry_price,
                                        'stop_loss': stop_loss,
                                        'risk_per_share': risk_per_share,
                                        'risk_pct': round(trade_risk_pct, 2),
                                        'target_1': target_1,
                                        't1_gain_pct': round((t1_gain / entry_price) * 100.0, 2),
                                        'target_2': target_2,
                                        't2_gain_pct': round((t2_gain / entry_price) * 100.0, 2),
                                        'rr_ratio': rr_ratio,
                                        'suggested_shares': suggested_shares,
                                        'total_investment': total_investment,
                                        'max_holding_days': MAX_HOLDING_DAYS,
                                        'rsi': round(float(rsi[i]), 1),
                                        'macd_hist': round(float(hist[i]), 2),
                                        'guidance': guidance,
                                        'g_score': g_score,
                                        'plan': plan,
                                        'tier': ('ULTRA' if g_score >= 2 else 'STANDARD')
                                    }

                # --- 2. RADAR WATCHLIST EVALUATION ---
                if is_uptrend and _pend is None:
                    dist_to_bb_low = (cmp - bb_low[i]) / cmp
                    if 0.0 <= dist_to_bb_low <= 0.035 and rsi[i] <= 45.0:
                        watchlist_candidates.append({
                            'symbol': meta['symbol'],
                            'name': meta['name'],
                            'industry': meta['industry'],
                            'cmp': cmp,
                            'bb_lower': round(bb_low[i], 2),
                            'dist_to_bb_low_pct': round(dist_to_bb_low * 100.0, 2),
                            'rsi': round(float(rsi[i]), 1),
                            'macd_hist': round(float(hist[i]), 2),
                            'macd_turning': "YES (Turning)" if hist[i] > hist[i-1] else "No (Waiting)",
                            'trigger_condition': "Consolidating at Lower BB Support; Watch for green hook"
                        })

                # --- 3. POSITION TRACKER (signals from last TRACK_DAYS, simulated to today) ---
                _ticker_busy = -1
                for offset in range(TRACK_DAYS, 0, -1):
                    p_i = i - offset
                    if p_i < 100: continue
                    if p_i <= _ticker_busy: continue
                    if not (c[p_i] > ema200[p_i] and ema50[p_i] > ema200[p_i]): continue
                    if ema50[p_i] < ema50[p_i-5] * 0.998: continue
                    _nm0 = _nifty_mom_on(dates[p_i].date())
                    if not (True if np.isnan(ret30[p_i]) else (ret30[p_i] >= _nm0 - 0.025)): continue
                    if bb_low_slope[p_i] < -0.012: continue
                    if not any(l[p_i-k] <= bb_low[p_i-k] * 1.005 for k in range(2)): continue
                    if not (any(rsi[p_i-k] <= 40 for k in range(3)) and rsi[p_i] > rsi[p_i-1] and rsi[p_i] >= 28): continue
                    if not (hist[p_i] > hist[p_i-1]): continue
                    _lw0 = min(o[p_i], c[p_i]) - l[p_i]
                    _wick0 = (_lw0 / (h[p_i] - l[p_i] + 1e-6) >= 0.20) or (c[p_i] >= (h[p_i] + l[p_i]) * 0.50) or (c[p_i] > c[p_i-1])
                    if not ((c[p_i] > bb_low[p_i]) and (c[p_i] > o[p_i]) and _wick0): continue

                    sig_close = c[p_i]
                    p_sl = round(min(l[p_i-2], l[p_i-1], l[p_i]) - (1.3 * atr[p_i]), 2)
                    p_risk = sig_close - p_sl
                    if p_risk <= 0 or not (0.012 <= (p_risk / sig_close) <= 0.08): continue

                    e_idx = p_i + 1
                    if e_idx > i: continue
                    e_price = float(o[e_idx])
                    if e_price <= p_sl: continue
                    if (e_price - sig_close) / sig_close > 0.015: continue
                    p_t1 = round(e_price * (1.0 + TARGET_1_GAIN_PCT), 2)
                    p_t2 = round(max(float(bb_up[p_i]), e_price * T2_MULT), 2)
                    _gs = 0
                    _nm = _nifty_mom_on(dates[p_i].date())
                    if (not np.isnan(ret30[p_i])) and ret30[p_i] > _nm + 0.03: _gs += 1
                    if rsi[p_i] < 55: _gs += 1
                    if hist[p_i] > 0: _gs += 1
                    if vol[p_i] > vol_sma[p_i]: _gs += 1
                    _tplan = 'FULL_RIDE' if _gs >= 3 else 'FULL_T1'
                    legs, _touched = simulate_exits(o, h, l, c, e_idx, i, e_price, p_sl, p_t1, p_t2, plan=_tplan)
                    tinfo = tracker_status(legs, e_price, float(c[i]), e_idx, dates, plan=_tplan, t1_touched=_touched)
                    recent_signals.append({
                        'symbol': meta['symbol'],
                        'name': meta['name'],
                        'industry': meta['industry'],
                        'trigger_date': dates[p_i].strftime('%d-%b-%Y'),
                        'days_ago': offset,
                        'entry_price': e_price,
                        'stop_loss': p_sl,
                        'target_1': p_t1,
                        'target_2': p_t2,
                        'cmp': float(c[i]),
                        'live_pnl_pct': tinfo['live_pnl_pct'],
                        'status': tinfo['status'],
                        'action': tinfo['action'],
                        'is_open': tinfo['is_open'],
                        'fresh_exit': tinfo['fresh_exit'],
                        'exit_date': tinfo['exit_date'],
                        'exit_reason': tinfo['exit_reason'],
                        'exit_price': tinfo['exit_price'],
                        'eff_sl': tinfo['eff_sl'],
                        't1_touched': tinfo['t1_touched'],
                        'days_held': tinfo['days_held'],
                        'guidance': ('RIDE' if _tplan == 'FULL_RIDE' else 'BOOK-FAST'),
                        'plan': _tplan,
                        'g_score': _gs,
                        'tier': ('ULTRA' if _gs >= 2 else 'STANDARD'),
                    })
                    _ticker_busy = legs[-1][0]
                if _pend is not None and _ticker_busy < i:
                    actionable_signals.append(_pend)
            except Exception:
                continue

    rank_order = {'A+': 3, 'A': 2, 'A-': 1}
    actionable_signals.sort(key=lambda x: (rank_order.get(x['rank'], 0), x['rr_ratio']), reverse=True)
    watchlist_candidates.sort(key=lambda x: x['dist_to_bb_low_pct'])
    recent_signals.sort(key=lambda x: x['days_ago'])

    return {
        'actionable': actionable_signals,
        'watchlist': watchlist_candidates,
        'recent': recent_signals
    }, regime

print("✅ Screener Batch Engine compiled with verified 7-fix constraints!")


# ============ EXECUTE SCAN ============

results, regime_info = run_screener(
    universe_df,
    capital=PORTFOLIO_CAPITAL,
    risk_pct=RISK_PER_TRADE_PCT
)

actionable = results['actionable']
watchlist = results['watchlist']
recent = results['recent']

print(f"\n{'='*85}")
print(f"  SCAN SUMMARY: {len(universe_df)} STOCKS SCANNED")
print(f"  - Actionable Buy Signals (Ready for Entry): {len(actionable)}")
print(f"  - Radar Watchlist (At Lower Band Support) : {len(watchlist)}")
print(f"  - Tracked Signals (Past {TRACK_DAYS} Days)      : {len(recent)}")
n_exit = sum(1 for t in recent if t.get("fresh_exit"))
print(f"  - Fresh EXIT alerts (T1/T2/SL/TIME)         : {n_exit}")
n_open = sum(1 for t in recent if t.get("is_open"))
print(f"  - Open tracked positions                : {n_open}")
print(f"{'='*85}\n")

if actionable:
    print("🟢 ACTIONABLE:")
    print(pd.DataFrame(actionable)[['rank','symbol','entry_price','stop_loss','target_1','target_2','rr_ratio','suggested_shares','total_investment']].to_string(index=False))
else:
    print("No stocks met 100% of the strict 7 confluence rules today.")
if watchlist:
    print("\n🟡 RADAR (top 25):")
    print(pd.DataFrame(watchlist[:25])[['symbol','cmp','bb_lower','dist_to_bb_low_pct','rsi','macd_hist','macd_turning']].to_string(index=False))


# ============ EXCEL REPORT ============
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, timedelta, timezone

# GitHub ka runner UTC pe chalta hai. Isliye har jagah India time force karo,
# warna 6:30 PM UTC ke baad report pe PICHLI date chhap jati hai.
IST = timezone(timedelta(hours=5, minutes=30))


def now_ist():
    """Hamesha Asia/Kolkata ka current time."""
    return datetime.now(IST)

def export_to_excel(results, regime_info, universe_label="NSE Top 1000"):
    wb = openpyxl.Workbook()

    font_title = Font(name='Calibri', size=15, bold=True, color='FFFFFF')
    font_subtitle = Font(name='Calibri', size=10, italic=True, color='CBD5E1')
    font_header = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
    font_body = Font(name='Calibri', size=10, color='1E293B')
    font_bold = Font(name='Calibri', size=10, bold=True, color='1E293B')
    font_green = Font(name='Calibri', size=10, bold=True, color='047857')
    font_red = Font(name='Calibri', size=10, bold=True, color='B91C1C')
    font_blue = Font(name='Calibri', size=10, bold=True, color='0369A1')

    fill_navy = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
    fill_slate = PatternFill(start_color='334155', end_color='334155', fill_type='solid')
    fill_header = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')
    fill_header_amber = PatternFill(start_color='78350F', end_color='78350F', fill_type='solid')
    fill_header_blue = PatternFill(start_color='0C4A6E', end_color='0C4A6E', fill_type='solid')
    fill_zebra = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    fill_light_green = PatternFill(start_color='D1FAE5', end_color='D1FAE5', fill_type='solid')
    fill_light_red = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')
    fill_cyan_bg = PatternFill(start_color='E0F2FE', end_color='E0F2FE', fill_type='solid')
    fill_amber_bg = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid')

    thin_border = Border(
        left=Side(border_style='thin', color='CBD5E1'),
        right=Side(border_style='thin', color='CBD5E1'),
        top=Side(border_style='thin', color='CBD5E1'),
        bottom=Side(border_style='thin', color='CBD5E1')
    )
    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')

    signals = results.get('actionable', [])
    watchlist = results.get('watchlist', [])
    recent = results.get('recent', [])

    # Sheet 1: Actionable Signals
    ws1 = wb.active
    ws1.title = "Actionable Swing Signals"
    ws1.views.sheetView[0].showGridLines = True
    ws1.merge_cells('A1:T1')
    ws1['A1'] = "BMR TRINITY SWING SCREENER - DAILY ACTIONABLE ORDERS"
    ws1['A1'].font = font_title; ws1['A1'].fill = fill_navy; ws1['A1'].alignment = align_center
    ws1.row_dimensions[1].height = 36

    today_str = now_ist().strftime("%d %b %Y, %I:%M %p")
    ws1.merge_cells('A2:T2')
    regime_txt = "BULLISH (Safe to Swing Trade)" if regime_info.get('is_bull_market') else "CORRECTION (Caution Recommended)"
    ws1['A2'] = f"Generated on: {today_str} | Universe: {universe_label} | Nifty 50 Status: {regime_txt} (CMP: ₹{regime_info.get('cmp', 0):,.2f})"
    ws1['A2'].font = font_subtitle; ws1['A2'].fill = fill_slate; ws1['A2'].alignment = align_center
    ws1.row_dimensions[2].height = 22

    headers_1 = [
        "Signal Rank", "NSE Ticker", "Company Name", "Industry / Sector",
        "CMP / Entry (₹)", "Stop Loss (₹)", "Risk / Share (₹)", "Risk (%)",
        "Target 1 (₹)", "T1 Return (%)", "Target 2 (₹)", "T2 Return (%)",
        "Reward:Risk", "Suggested Qty", "Investment (₹)", "Max Holding",
        "RSI (14)", "MACD Hist", "Guidance", "Tier"
    ]
    for c_i, h in enumerate(headers_1, 1):
        cell = ws1.cell(row=4, column=c_i, value=h)
        cell.font = font_header; cell.fill = fill_header; cell.alignment = align_center; cell.border = thin_border
    ws1.row_dimensions[4].height = 26

    r_idx = 4
    if not signals:
        r_idx += 1
        ws1.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=len(headers_1))
        c = ws1.cell(row=r_idx, column=1, value="No stocks met 100% of strict 7 confluence rules today. (Check 'Radar Watchlist' tab for stocks setting up near lower band).")
        c.font = font_bold; c.alignment = align_center
        ws1.row_dimensions[r_idx].height = 28
    else:
        for idx, sig in enumerate(signals):
            r_idx += 1
            row_fill = fill_zebra if (idx % 2 == 0) else PatternFill(fill_type=None)
            vals = [
                sig['rank'], sig['symbol'], sig['name'], sig['industry'],
                sig['entry_price'], sig['stop_loss'], sig['risk_per_share'], sig['risk_pct'] / 100.0,
                sig['target_1'], sig['t1_gain_pct'] / 100.0, sig['target_2'], sig['t2_gain_pct'] / 100.0,
                sig['rr_ratio'], sig['suggested_shares'], sig['total_investment'],
                f"{sig['max_holding_days']} Days", sig['rsi'], sig['macd_hist'], ("RIDE" if sig['guidance']=="RIDE" else "BOOK-FAST"), ("ULTRA" if sig['tier']=="ULTRA" else "STANDARD")
            ]
            for col_i, val in enumerate(vals, 1):
                cell = ws1.cell(row=r_idx, column=col_i, value=val)
                cell.font = font_body; cell.border = thin_border; cell.fill = row_fill
                if col_i in [1, 2, 16, 19, 20]: cell.alignment = align_center
                elif col_i in [3, 4]: cell.alignment = align_left
                else: cell.alignment = align_right

                if col_i == 1:
                    if val == 'A+': cell.fill = fill_light_green; cell.font = font_green
                    elif val == 'A': cell.fill = fill_cyan_bg; cell.font = font_bold
                elif col_i in [5, 6, 7, 9, 11, 15]:
                    cell.number_format = '₹#,##0.00'
                    if col_i == 6: cell.font = font_red
                elif col_i in [8, 10, 12]:
                    cell.number_format = '+0.0%;-0.0%;0.0%'
                    if col_i in [10, 12]: cell.font = font_green
                    elif col_i == 8: cell.font = font_red
                elif col_i == 13: cell.number_format = '0.00'
                elif col_i == 14: cell.number_format = '#,##0'
                elif col_i == 19:
                    cell.alignment = align_center
                    if val == 'RIDE': cell.fill = fill_light_green; cell.font = font_green
                    else: cell.fill = fill_amber_bg; cell.font = font_bold
                elif col_i == 20:
                    cell.alignment = align_center
                    if val == 'ULTRA': cell.fill = fill_amber_bg; cell.font = font_bold
            ws1.row_dimensions[r_idx].height = 20

    # Sheet 2: Radar Watchlist
    ws2 = wb.create_sheet(title="Radar Watchlist (Near Trigger)")
    ws2.views.sheetView[0].showGridLines = True
    ws2.merge_cells('A1:J1')
    ws2['A1'] = "HIGH-PRIORITY RADAR WATCHLIST (TESTING LOWER BB SUPPORT IN UPTREND)"
    ws2['A1'].font = font_title; ws2['A1'].fill = fill_header_amber; ws2['A1'].alignment = align_center
    ws2.row_dimensions[1].height = 36
    ws2.merge_cells('A2:J2')
    ws2['A2'] = "Criteria: Primary Uptrend (>200 EMA & 50>200 EMA) + Trading within 3.5% of Lower Bollinger Band + RSI < 45. Monitor for green hook."
    ws2['A2'].font = font_subtitle; ws2['A2'].fill = fill_slate; ws2['A2'].alignment = align_center
    ws2.row_dimensions[2].height = 22
    headers_2 = ["NSE Ticker", "Company Name", "Sector / Industry", "CMP (₹)", "Lower BB (₹)", "Dist to Lower BB (%)", "RSI (14)", "MACD Hist", "MACD Turning?", "Trigger Note"]
    for c_i, h in enumerate(headers_2, 1):
        cell = ws2.cell(row=4, column=c_i, value=h)
        cell.font = font_header; cell.fill = fill_slate; cell.alignment = align_center; cell.border = thin_border
    ws2.row_dimensions[4].height = 26
    r_idx = 4
    for idx, item in enumerate(watchlist[:50]):
        r_idx += 1
        row_fill = fill_zebra if (idx % 2 == 0) else PatternFill(fill_type=None)
        w_vals = [
            item['symbol'], item['name'], item['industry'], item['cmp'], item['bb_lower'],
            item['dist_to_bb_low_pct'] / 100.0, item['rsi'], item['macd_hist'], item['macd_turning'], item['trigger_condition']
        ]
        for col_i, val in enumerate(w_vals, 1):
            cell = ws2.cell(row=r_idx, column=col_i, value=val)
            cell.font = font_body; cell.border = thin_border; cell.fill = row_fill
            if col_i in [1, 9]: cell.alignment = align_center
            elif col_i in [2, 3, 10]: cell.alignment = align_left
            else: cell.alignment = align_right
            if col_i in [4, 5]: cell.number_format = '₹#,##0.00'
            elif col_i == 6: cell.number_format = '0.0%'
        ws2.row_dimensions[r_idx].height = 20

    # Sheet 3: Position Tracker (25-day, live simulation)
    ws3 = wb.create_sheet(title="Position Tracker (25 Days)")
    ws3.views.sheetView[0].showGridLines = True
    ws3.merge_cells('A1:M1')
    ws3['A1'] = "POSITION TRACKER - OPEN HOLDINGS + EXIT ALERTS (AUTO-SIMULATED TO TODAY)"
    ws3['A1'].font = font_title; ws3['A1'].fill = fill_header_blue; ws3['A1'].alignment = align_center
    ws3.row_dimensions[1].height = 36
    ws3.merge_cells('A2:M2')
    ws3['A2'] = "M4: FULL_T1 = T1 par 100% exit | FULL_RIDE = T1 par SL cost par, T2 par exit | SL = exit now | TIME = 25 sessions over, exit at CMP"
    ws3['A2'].font = font_subtitle; ws3['A2'].fill = fill_slate; ws3['A2'].alignment = align_center
    ws3.row_dimensions[2].height = 22
    headers_3 = ["Symbol", "Trigger Date", "Days Held", "Entry (₹)", "SL (₹)", "T1 (₹)", "T2 (₹)", "CMP (₹)", "Live P&L (%)", "Status", "Tier", "Action", "Exit (₹)"]
    for c_i, h in enumerate(headers_3, 1):
        cell = ws3.cell(row=4, column=c_i, value=h)
        cell.font = font_header; cell.fill = fill_slate; cell.alignment = align_center; cell.border = thin_border
    ws3.row_dimensions[4].height = 26
    tr_open = sorted([t for t in recent if t.get('is_open')], key=lambda x: x['days_held'], reverse=True)
    tr_closed = sorted([t for t in recent if not t.get('is_open')], key=lambda x: (not x.get('fresh_exit'), x['days_ago']))
    r_idx = 4
    for idx, item in enumerate((tr_open + tr_closed)[:50]):
        r_idx += 1
        row_fill = fill_zebra if (idx % 2 == 0) else PatternFill(fill_type=None)
        row_vals = [item['symbol'], item['trigger_date'], item['days_held'], item['entry_price'], (item.get('eff_sl') if item.get('eff_sl') is not None else item['stop_loss']), item['target_1'], item['target_2'], item['cmp'], item['live_pnl_pct'] / 100.0, item['status'], ("ULTRA" if item.get('tier')=="ULTRA" else "STANDARD"), item['action'], (item.get('exit_price') if item.get('exit_price') is not None else item['cmp'])]
        for col_i, val in enumerate(row_vals, 1):
            cell = ws3.cell(row=r_idx, column=col_i, value=val)
            cell.font = font_body; cell.border = thin_border; cell.fill = row_fill
            if col_i in [1, 2, 3, 10, 11]: cell.alignment = align_center
            elif col_i == 12: cell.alignment = align_left
            else: cell.alignment = align_right
            if col_i in [4, 5, 6, 7, 8, 13]: cell.number_format = '₹#,##0.00'
            elif col_i == 9:
                cell.number_format = '+0.0%;-0.0%;0.0%'
                cell.font = font_green if val >= 0 else font_red
            elif col_i == 10:
                if "T2 DONE" in val or "T1 FULL" in val: cell.fill = fill_light_green; cell.font = font_green
                elif "SL HIT" in val: cell.fill = fill_light_red; cell.font = font_red
                elif "TIME UP" in val: cell.fill = fill_amber_bg; cell.font = font_bold
                elif "OPEN" in val: cell.fill = fill_cyan_bg; cell.font = font_blue
            elif col_i == 11:
                cell.alignment = align_center
                if val == 'ULTRA': cell.fill = fill_amber_bg; cell.font = font_bold
        ws3.row_dimensions[r_idx].height = 20
    if r_idx == 4:
        ws3.merge_cells('A5:M5')
        ws3['A5'] = "No tracked signals in last 25 days."
        ws3['A5'].font = font_bold; ws3['A5'].alignment = align_center

    # Sheet 4: Regime
    ws4 = wb.create_sheet(title="Market Regime & Breadth")
    ws4.views.sheetView[0].showGridLines = True
    ws4.merge_cells('A1:E1')
    ws4['A1'] = "NIFTY 50 MACRO REGIME & UNIVERSE BREADTH"
    ws4['A1'].font = font_title; ws4['A1'].fill = fill_navy; ws4['A1'].alignment = align_center
    ws4.row_dimensions[1].height = 36
    ws4.cell(row=3, column=1, value="Metric").font = font_header; ws4.cell(row=3, column=1).fill = fill_slate
    ws4.cell(row=3, column=2, value="Current Value").font = font_header; ws4.cell(row=3, column=2).fill = fill_slate
    ws4.cell(row=3, column=3, value="Operational Meaning").font = font_header; ws4.cell(row=3, column=3).fill = fill_slate
    nifty_rows = [
        ("NIFTY 50 Current Market Price (CMP)", f"₹{regime_info.get('cmp', 0):,.2f}", "Benchmark reference price"),
        ("NIFTY 50 50-Day EMA", f"₹{regime_info.get('ema50', 0):,.2f}", "Intermediate trend benchmark"),
        ("NIFTY 50 200-Day EMA", f"₹{regime_info.get('ema200', 0):,.2f}", "Long-term secular bull/bear dividing line"),
        ("Market Regime Status", "BULL MARKET" if regime_info.get('is_bull_market') else "CORRECTION / CHOPPY", "Trade full size in Bull; Reduce risk in Correction"),
        ("NIFTY 30-Day Momentum Return", f"{regime_info.get('nifty_30d_return', 0)*100:+.2f}%", "Stocks must beat this benchmark to qualify (RS Filter)")
    ]
    for r_i, (k, v, desc) in enumerate(nifty_rows, 4):
        c1 = ws4.cell(row=r_i, column=1, value=k); c2 = ws4.cell(row=r_i, column=2, value=v); c3 = ws4.cell(row=r_i, column=3, value=desc)
        c1.font = font_bold; c1.border = thin_border; c2.font = font_body; c2.border = thin_border; c3.font = font_body; c3.border = thin_border
        if "BULL MARKET" in v: c2.font = font_green; c2.fill = fill_light_green
        elif "CORRECTION" in v: c2.font = font_red; c2.fill = fill_light_red
        ws4.row_dimensions[r_i].height = 22

    # Sheet 5: SOP
    ws5 = wb.create_sheet(title="Strategy Execution SOP")
    ws5.views.sheetView[0].showGridLines = True
    ws5.merge_cells('A1:C1')
    ws5['A1'] = "BMR TRINITY SWING STRATEGY - STEP-BY-STEP EXECUTION SOP"
    ws5['A1'].font = font_title; ws5['A1'].fill = fill_navy; ws5['A1'].alignment = align_center
    ws5.row_dimensions[1].height = 36
    sop_steps = [
        ("Step 1: Daily Screener Run Time", "Run screener daily between 3:15 PM - 3:25 PM (to catch closing prices) or after market hours (7:00 PM - 11:00 PM)."),
        ("Step 2: Order Placement (Subah 9:15 AM)", "Agley din market open hote hi 'Suggested Qty' ke hisab se order execute karein. Do not chase if stock opens > 1.5% gap-up."),
        ("Step 3: Immediate Stop Loss Order", "Broker app (Zerodha / Groww / Angel One / Upstox) me sheet ka exact 'Stop Loss (₹)' ka GTT / SL-Limit order turant laga dein."),
        ("Step 4: Target 1 - M4 Rule", "Plan FULL_T1 (score 0-2): T1 hit hote hi 100% SELL karo. Plan FULL_RIDE (score 3-4): KUCH MAT BECHO - sirf SL ko Cost + 0.5% par trail karo."),
        ("Step 5: Target 2 / Trailed SL (RIDE only)", "FULL_RIDE position T2 (+6% ya Upper BB) par 100% exit karo. T2 se pehle trailed SL (Cost + 0.5%) hite to wahi niklo - chhota profit lock."),
        ("Step 6: Time Stop Discipline", "Agar stock 25 trading sessions (lagbhag 1 se 1.2 mahine) me Target 2 hit na kare, toh Day 26 par market price par position band kar dein.")
    ]
    ws5.cell(row=3, column=1, value="SOP Step").font = font_header; ws5.cell(row=3, column=1).fill = fill_slate
    ws5.cell(row=3, column=2, value="Execution Instructions (Hindi / English)").font = font_header; ws5.cell(row=3, column=2).fill = fill_slate
    for r_i, (step, desc) in enumerate(sop_steps, 4):
        c1 = ws5.cell(row=r_i, column=1, value=step); c2 = ws5.cell(row=r_i, column=2, value=desc)
        c1.font = font_bold; c1.border = thin_border; c2.font = font_body; c2.border = thin_border
        ws5.row_dimensions[r_i].height = 24

    for ws in [ws1, ws2, ws3, ws4, ws5]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row in [1, 2]: continue
                val_str = str(cell.value or '')
                if len(val_str) > max_len: max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

    filename = f"BMR_Trinity_NSE_Screener_Report_{now_ist().strftime('%Y%m%d_%H%M')}.xlsx"
    wb.save(filename)
    return filename

excel_file = export_to_excel(results, regime_info, universe_label=f"NSE {UNIVERSE.upper()}")
print(f"✅ Excel report created successfully: {excel_file}")

# ============ STUNNING HTML EMAIL REPORT (v2: tracker + guidance) ============
import json
import smtplib
from html import escape as _esc
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.application import MIMEApplication


def E(x):
    return _esc(str(x))


def build_html_report(results, regime_info, universe_label):
    actionable = results.get('actionable', [])
    watchlist = results.get('watchlist', [])
    tracked = results.get('recent', [])
    fresh = [t for t in tracked if t.get('fresh_exit')]
    openpos = [t for t in tracked if t.get('is_open')]
    ultra = [s for s in actionable if s.get('tier') == 'ULTRA']
    std = [s for s in actionable if s.get('tier') != 'ULTRA']
    date_str = now_ist().strftime("%A, %d %B %Y")
    # Jis waqt report actually bani, wahi dikhao (pehle "6:00 PM IST" hardcoded tha)
    sent_str = now_ist().strftime("%I:%M %p IST").lstrip("0")

    is_bull = regime_info.get('is_bull_market', False)
    regime_txt = "BULL MARKET" if is_bull else "CORRECTION"
    regime_color = "#16C784" if is_bull else "#EA3943"
    regime_bg = "#0B2E23" if is_bull else "#3A1215"

    # ---- 1. EXIT ALERTS ----
    ex_rows = ""
    for t in fresh:
        pnl_c = "#16C784" if t['live_pnl_pct'] >= 0 else "#EA3943"
        if "T2 DONE" in t['status']:
            badge = '<span style="font-size:11px;color:#06281C;font-weight:800;background-color:#16C784;border-radius:4px;padding:3px 8px;">EXIT: T2 HIT</span>'
        elif "SL HIT" in t['status']:
            badge = '<span style="font-size:11px;color:#fff;font-weight:800;background-color:#EA3943;border-radius:4px;padding:3px 8px;">EXIT: SL HIT</span>'
        elif "T1 FULL" in t['status']:
            badge = '<span style="font-size:11px;color:#06281C;font-weight:800;background-color:#16C784;border-radius:4px;padding:3px 8px;">EXIT: T1 FULL</span>'
        elif "TIME UP" in t['status']:
            badge = '<span style="font-size:11px;color:#3A2E00;font-weight:800;background-color:#F0B90B;border-radius:4px;padding:3px 8px;">EXIT: TIME UP</span>'
        else:
            badge = '<span style="font-size:11px;color:#06283D;font-weight:800;background-color:#38BDF8;border-radius:4px;padding:3px 8px;">EXIT</span>'
        ex_rows += f"""<tr style="background-color:#141B2B;">
      <td style="font-size:13px;color:#FFFFFF;font-weight:bold;padding:10px 6px 10px 12px;border-top:1px solid #232C42;">{('💎' if t.get('tier')=='ULTRA' else '📋')} {E(t['symbol'])}<div style="font-size:10px;color:#8A94A6;font-weight:normal;">{E(t['exit_date'])} • {t['days_held']}d held</div></td>
      <td align="right" style="font-size:13px;color:#FFFFFF;padding:10px 6px;border-top:1px solid #232C42;">₹{t['entry_price']:,.2f}</td>
      <td align="right" style="font-size:13px;color:{pnl_c};font-weight:bold;padding:10px 6px;border-top:1px solid #232C42;">{t['live_pnl_pct']:+.1f}%</td>
      <td align="right" style="padding:10px 12px 10px 6px;border-top:1px solid #232C42;">{badge}</td>
    </tr>
    <tr><td colspan="4" style="font-size:11px;color:#8A94A6;padding:0 6px 2px 12px;">Exit @ {("₹{:,.2f}".format(t["exit_price"]) if t.get("exit_price") is not None else "-")} &nbsp;•&nbsp; 🛑 SL ₹{t["stop_loss"]:,.2f} &nbsp;•&nbsp; 🎯 T1 ₹{t["target_1"]:,.2f} &nbsp;•&nbsp; 🎯 T2 ₹{t["target_2"]:,.2f}</td></tr>
    <tr><td colspan="4" style="font-size:11px;color:#F0B90B;padding:0 6px 10px 12px;">👉 {E(t['action'])}</td></tr>"""
    if not ex_rows:
        ex_rows = """<tr><td style="font-size:13px;color:#16C784;text-align:center;padding:14px;">👍 Koi fresh exit nahi — sab positions track par hain.</td></tr>"""

    # ---- 2. OPEN POSITIONS ----
    op_rows = ""
    for t in openpos[:15]:
        pnl_c = "#16C784" if t['live_pnl_pct'] >= 0 else "#EA3943"
        t1tag = "T1✅ " if t.get('t1_touched') else ""
        op_rows += f"""<tr>
      <td style="font-size:13px;color:#FFFFFF;font-weight:bold;padding:10px 6px 10px 12px;border-top:1px solid #232C42;">{('💎' if t.get('tier')=='ULTRA' else '📋')} {E(t['symbol'])}<div style="font-size:10px;color:#8A94A6;font-weight:normal;">{t1tag}{E(t['trigger_date'])} • {t['days_held']}d</div></td>
      <td align="right" style="font-size:13px;color:#FFFFFF;padding:10px 6px;border-top:1px solid #232C42;">₹{t['entry_price']:,.2f}</td>
      <td align="right" style="font-size:13px;color:#FFFFFF;padding:10px 6px;border-top:1px solid #232C42;">₹{t['cmp']:,.2f}</td>
      <td align="right" style="font-size:13px;color:{pnl_c};font-weight:bold;padding:10px 6px;border-top:1px solid #232C42;">{t['live_pnl_pct']:+.1f}%</td>
      <td style="font-size:11px;color:#B0B8C8;padding:10px 12px 10px 6px;border-top:1px solid #232C42;">{E(t['action'])}</td>
    </tr>
    <tr><td colspan="5" style="font-size:11px;color:#8A94A6;padding:0 6px 10px 12px;">🛑 SL ₹{(t["eff_sl"] if t.get("eff_sl") is not None else t["stop_loss"]):,.2f}{" (trailed🔒)" if t.get("eff_sl") is not None else ""} &nbsp;•&nbsp; 🎯 T1 ₹{t["target_1"]:,.2f} &nbsp;•&nbsp; 🎯 T2 ₹{t["target_2"]:,.2f}</td></tr>"""
    if not op_rows:
        op_rows = """<tr><td colspan="5" style="font-size:13px;color:#8A94A6;text-align:center;padding:14px;">Koi open tracked position nahi.</td></tr>"""

    # ---- 3. NEW ACTIONABLE (with guidance) ----
    rank_bg = {'A+': '#16C784', 'A': '#38BDF8', 'A-': '#F0B90B'}
    act_rows = ""
    for i, s in enumerate(ultra[:20]):
        bg = ' style="background-color:#141B2B;"' if i % 2 == 1 else ''
        rb = rank_bg.get(s['rank'], '#F0B90B')
        if s.get('guidance') == 'RIDE':
            g = '<span style="font-size:11px;color:#06281C;font-weight:800;background-color:#16C784;border-radius:4px;padding:3px 8px;">🚀 FULL-RIDE</span>'
        else:
            g = '<span style="font-size:11px;color:#3A2E00;font-weight:800;background-color:#F0B90B;border-radius:4px;padding:3px 8px;">⚡ FULL-EXIT</span>'
        act_rows += f"""<tr{bg}>
      <td style="font-size:13px;color:#FFFFFF;font-weight:bold;padding:10px 6px 10px 12px;border-top:1px solid #232C42;">{E(s['symbol'])}<div style="font-size:10px;color:#8A94A6;font-weight:normal;">{E(s['industry'][:22])}</div></td>
      <td align="center" style="padding:10px 6px;border-top:1px solid #232C42;"><span style="font-size:11px;color:#06281C;font-weight:800;background-color:{rb};border-radius:4px;padding:3px 8px;">{E(s['rank'])}</span></td>
      <td align="right" style="font-size:13px;color:#FFFFFF;padding:10px 6px;border-top:1px solid #232C42;">₹{s['entry_price']:,.2f}</td>
      <td align="right" style="font-size:13px;color:#EA3943;padding:10px 6px;border-top:1px solid #232C42;">₹{s['stop_loss']:,.2f}</td>
      <td align="right" style="font-size:13px;color:#16C784;padding:10px 6px;border-top:1px solid #232C42;">₹{s['target_1']:,.2f}</td>
      <td align="right" style="font-size:13px;color:#16C784;padding:10px 6px;border-top:1px solid #232C42;">₹{s['target_2']:,.2f}</td>
      <td align="center" style="padding:10px 12px 10px 6px;border-top:1px solid #232C42;">{g}</td>
    </tr>"""
    if not act_rows:
        act_rows = """<tr><td colspan="7" style="font-size:13px;color:#8A94A6;text-align:center;padding:16px;border-top:1px solid #232C42;">No 💎 ultra signals today.</td></tr>"""

    # ---- 3b. STANDARD ACTIONABLE ----
    std_rows = ""
    for i, s in enumerate(std[:20]):
        bg = ' style="background-color:#141B2B;"' if i % 2 == 1 else ''
        rb = rank_bg.get(s['rank'], '#F0B90B')
        g = '<span style="font-size:11px;color:#3A2E00;font-weight:800;background-color:#F0B90B;border-radius:4px;padding:3px 8px;">⚡ FULL-EXIT</span>'
        std_rows += f"""<tr{bg}>
      <td style="font-size:13px;color:#FFFFFF;font-weight:bold;padding:10px 6px 10px 12px;border-top:1px solid #232C42;">{E(s['symbol'])}<div style="font-size:10px;color:#8A94A6;font-weight:normal;">{E(s['industry'][:22])}</div></td>
      <td align="center" style="padding:10px 6px;border-top:1px solid #232C42;"><span style="font-size:11px;color:#06281C;font-weight:800;background-color:{rb};border-radius:4px;padding:3px 8px;">{E(s['rank'])}</span></td>
      <td align="right" style="font-size:13px;color:#FFFFFF;padding:10px 6px;border-top:1px solid #232C42;">₹{s['entry_price']:,.2f}</td>
      <td align="right" style="font-size:13px;color:#EA3943;padding:10px 6px;border-top:1px solid #232C42;">₹{s['stop_loss']:,.2f}</td>
      <td align="right" style="font-size:13px;color:#16C784;padding:10px 6px;border-top:1px solid #232C42;">₹{s['target_1']:,.2f}</td>
      <td align="right" style="font-size:13px;color:#16C784;padding:10px 6px;border-top:1px solid #232C42;">₹{s['target_2']:,.2f}</td>
      <td align="center" style="padding:10px 12px 10px 6px;border-top:1px solid #232C42;">{g}</td>
    </tr>"""
    if not std_rows:
        std_rows = """<tr><td colspan="7" style="font-size:13px;color:#8A94A6;text-align:center;padding:16px;border-top:1px solid #232C42;">No standard signals today.</td></tr>"""

    # ---- 4. RADAR ----
    wl_rows = ""
    for i, w in enumerate(watchlist[:12]):
        bg = ' style="background-color:#141B2B;"' if i % 2 == 1 else ''
        turning = "🟢 YES" if "YES" in w['macd_turning'] else "⏳ Wait"
        wl_rows += f"""<tr{bg}>
      <td style="font-size:13px;color:#FFFFFF;font-weight:bold;padding:10px 6px 10px 12px;border-top:1px solid #232C42;">{E(w['symbol'])}</td>
      <td align="right" style="font-size:13px;color:#FFFFFF;padding:10px 6px;border-top:1px solid #232C42;">₹{w['cmp']:,.2f}</td>
      <td align="right" style="font-size:13px;color:#F0B90B;padding:10px 6px;border-top:1px solid #232C42;">{w['dist_to_bb_low_pct']:.1f}%</td>
      <td align="right" style="font-size:13px;color:#FFFFFF;padding:10px 6px;border-top:1px solid #232C42;">{w['rsi']:.1f}</td>
      <td align="right" style="font-size:12px;color:#B0B8C8;padding:10px 12px 10px 6px;border-top:1px solid #232C42;">{turning}</td>
    </tr>"""
    if not wl_rows:
        wl_rows = """<tr><td colspan="5" style="font-size:13px;color:#8A94A6;text-align:center;padding:16px;border-top:1px solid #232C42;">Radar clear — nothing near lower band today.</td></tr>"""

    html = f"""<!DOCTYPE html><html><head><meta charset="utf-8"></head>
<body style="margin:0;padding:0;background-color:#0B0E14;font-family:'Segoe UI',Arial,Helvetica,sans-serif;">
<table width="100%" cellpadding="0" cellspacing="0" style="background-color:#0B0E14;padding:24px 12px;"><tr><td align="center">
<table width="640" cellpadding="0" cellspacing="0" style="max-width:640px;">
<tr><td height="4" style="background-color:#F0B90B;border-radius:8px 8px 0 0;font-size:0;">&nbsp;</td></tr>
<tr><td style="background-color:#111827;padding:28px 26px 20px 26px;border-left:1px solid #232C42;border-right:1px solid #232C42;">
  <div style="font-size:12px;letter-spacing:3px;color:#F0B90B;font-weight:bold;">● BMR TRINITY SWING SCREENER</div>
  <div style="font-size:28px;font-weight:800;color:#FFFFFF;margin:8px 0 4px 0;">Daily Action Report</div>
  <div style="font-size:13px;color:#8A94A6;">{date_str} &nbsp;•&nbsp; {E(universe_label)} &nbsp;•&nbsp; {sent_str}</div>
  <table width="100%" cellpadding="0" cellspacing="0" style="margin-top:16px;"><tr>
    <td width="32%" align="center" style="background-color:{regime_bg};border:1px solid {regime_color};border-radius:8px;padding:10px;">
      <div style="font-size:10px;color:#8A94A6;letter-spacing:1px;">NIFTY REGIME</div>
      <div style="font-size:14px;color:{regime_color};font-weight:800;margin-top:2px;">{regime_txt}</div></td>
    <td width="2%"></td>
    <td width="32%" align="center" style="background-color:#151B29;border:1px solid #232C42;border-radius:8px;padding:10px;">
      <div style="font-size:10px;color:#8A94A6;letter-spacing:1px;">NIFTY 50 CMP</div>
      <div style="font-size:14px;color:#FFFFFF;font-weight:800;margin-top:2px;">₹{regime_info.get('cmp', 0):,.2f}</div></td>
    <td width="2%"></td>
    <td width="32%" align="center" style="background-color:#151B29;border:1px solid #232C42;border-radius:8px;padding:10px;">
      <div style="font-size:10px;color:#8A94A6;letter-spacing:1px;">NIFTY 30D</div>
      <div style="font-size:14px;color:#FFFFFF;font-weight:800;margin-top:2px;">{regime_info.get('nifty_30d_return', 0)*100:+.2f}%</div></td>
  </tr></table>
  <table width="100%" cellpadding="0" cellspacing="0" style="margin-top:10px;"><tr>
    <td width="24%" align="center" style="background-color:#3A1215;border:1px solid #EA3943;border-radius:8px;padding:10px;">
      <div style="font-size:20px;font-weight:800;color:#EA3943;">{len(fresh)}</div><div style="font-size:10px;color:#8A94A6;letter-spacing:1px;">EXIT NOW</div></td>
    <td width="1%"></td>
    <td width="24%" align="center" style="background-color:#0B2440;border:1px solid #38BDF8;border-radius:8px;padding:10px;">
      <div style="font-size:20px;font-weight:800;color:#38BDF8;">{len(openpos)}</div><div style="font-size:10px;color:#8A94A6;letter-spacing:1px;">OPEN</div></td>
    <td width="1%"></td>
    <td width="24%" align="center" style="background-color:#0B2E23;border:1px solid #16C784;border-radius:8px;padding:10px;">
      <div style="font-size:20px;font-weight:800;color:#16C784;">{len(actionable)}</div><div style="font-size:10px;color:#8A94A6;letter-spacing:1px;">NEW</div></td>
    <td width="1%"></td>
    <td width="24%" align="center" style="background-color:#2E230B;border:1px solid #F0B90B;border-radius:8px;padding:10px;">
      <div style="font-size:20px;font-weight:800;color:#F0B90B;">{len(watchlist)}</div><div style="font-size:10px;color:#8A94A6;letter-spacing:1px;">RADAR</div></td>
  </tr></table>
</td></tr>
<tr><td style="background-color:#111827;padding:0 26px 20px 26px;border-left:1px solid #232C42;border-right:1px solid #232C42;">
  <div style="font-size:11px;letter-spacing:2px;color:#EA3943;font-weight:bold;margin-bottom:10px;">🚨 EXIT ALERTS ({len(fresh)}) — TURANT ACTION LO</div>
  <table width="100%" cellpadding="0" cellspacing="0" style="border:1px solid #EA3943;border-radius:8px;">
  <tr style="background-color:#1A2233;">
    <td style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px 10px 12px;">STOCK</td>
    <td align="right" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px;">ENTRY</td>
    <td align="right" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px;">P&amp;L</td>
    <td align="right" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 12px 10px 6px;">SIGNAL</td>
  </tr>{ex_rows}</table>
</td></tr>
<tr><td style="background-color:#111827;padding:0 26px 20px 26px;border-left:1px solid #232C42;border-right:1px solid #232C42;">
  <div style="font-size:11px;letter-spacing:2px;color:#8A94A6;font-weight:bold;margin-bottom:10px;">📌 OPEN POSITIONS ({len(openpos)} SHOWN {min(len(openpos), 15)})</div>
  <table width="100%" cellpadding="0" cellspacing="0" style="border:1px solid #232C42;border-radius:8px;">
  <tr style="background-color:#1A2233;">
    <td style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px 10px 12px;">STOCK</td>
    <td align="right" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px;">ENTRY</td>
    <td align="right" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px;">CMP</td>
    <td align="right" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px;">P&amp;L</td>
    <td style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 12px 10px 6px;">GUIDANCE</td>
  </tr>{op_rows}</table>
</td></tr>
<tr><td style="background-color:#111827;padding:0 26px 20px 26px;border-left:1px solid #232C42;border-right:1px solid #232C42;">
  <div style="font-size:11px;letter-spacing:2px;color:#8A94A6;font-weight:bold;margin-bottom:10px;">💎 ULTRA QUALITY ({len(ultra)}) — Backtest: 85% WR • PF 1.8 — best signals, normal size</div>
  <table width="100%" cellpadding="0" cellspacing="0" style="border:1px solid #232C42;border-radius:8px;">
  <tr style="background-color:#1A2233;">
    <td style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px 10px 12px;">STOCK</td>
    <td align="center" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px;">RANK</td>
    <td align="right" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px;">ENTRY</td>
    <td align="right" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px;">SL</td>
    <td align="right" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px;">T1</td>
    <td align="right" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px;">T2</td>
    <td align="center" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 12px 10px 6px;">PLAN</td>
  </tr>{act_rows}</table>
</td></tr>
<tr><td style="background-color:#111827;padding:0 26px 20px 26px;border-left:1px solid #232C42;border-right:1px solid #232C42;">
  <div style="font-size:11px;letter-spacing:2px;color:#8A94A6;font-weight:bold;margin-bottom:10px;">📋 STANDARD ({len(std)}) — Backtest: 78% WR • PF 1.0 — chhota size, T1 par poora niklo</div>
  <table width="100%" cellpadding="0" cellspacing="0" style="border:1px solid #232C42;border-radius:8px;">
  <tr style="background-color:#1A2233;">
    <td style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px 10px 12px;">STOCK</td>
    <td align="center" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px;">RANK</td>
    <td align="right" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px;">ENTRY</td>
    <td align="right" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px;">SL</td>
    <td align="right" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px;">T1</td>
    <td align="right" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px;">T2</td>
    <td align="center" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 12px 10px 6px;">PLAN</td>
  </tr>{std_rows}</table>
</td></tr>
<tr><td style="background-color:#111827;padding:0 26px 20px 26px;border-left:1px solid #232C42;border-right:1px solid #232C42;">
  <div style="font-size:11px;letter-spacing:2px;color:#8A94A6;font-weight:bold;margin-bottom:10px;">🟡 RADAR WATCHLIST (TOP {min(len(watchlist), 12)} OF {len(watchlist)})</div>
  <table width="100%" cellpadding="0" cellspacing="0" style="border:1px solid #232C42;border-radius:8px;">
  <tr style="background-color:#1A2233;">
    <td style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px 10px 12px;">STOCK</td>
    <td align="right" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px;">CMP</td>
    <td align="right" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px;">DIST BB↓</td>
    <td align="right" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 6px;">RSI</td>
    <td align="right" style="font-size:10px;color:#8A94A6;font-weight:bold;padding:10px 12px 10px 6px;">MACD</td>
  </tr>{wl_rows}</table>
</td></tr>
<tr><td style="background-color:#0D1220;padding:18px 26px;border:1px solid #232C42;border-radius:0 0 8px 8px;">
  <div style="font-size:11px;color:#5A6579;line-height:1.6;">📎 Full Excel (5 sheets incl. Position Tracker) attached &nbsp;•&nbsp; 🚀 FULL-RIDE = T1 par sirf SL cost par, bechna kuch nahi &nbsp;•&nbsp; ⚡ FULL-EXIT = T1 par 100% book karo &nbsp;•&nbsp; 📌 Entry/SL/T1/T2 aaj ke close par based hain, kal ke open par thoda adjust ho sakte hain</div>
  <div style="font-size:11px;color:#5A6579;margin-top:8px;">⚠️ Educational purpose only, investment advice nahi hai.</div>
</td></tr>
</table></td></tr></table></body></html>"""
    return html


# ============ BUILD + SAVE + SEND ============
print("\n[*] Building stunning HTML report...")
html_report = build_html_report(results, regime_info, universe_label=f"NSE {UNIVERSE.upper()}")
with open("report.html", "w", encoding="utf-8") as f:
    f.write(html_report)
print("[+] report.html saved")
with open("results.json", "w", encoding="utf-8") as f:
    json.dump({"actionable": results.get("actionable", []), "watchlist": results.get("watchlist", []),
               "recent": results.get("recent", []), "regime": regime_info}, f, indent=1, default=str)
print("[+] results.json saved")

email_from = os.environ.get("MY_EMAIL")
app_password = os.environ.get("MY_APP_PASSWORD")

if email_from and app_password:
    print("[*] Sending email with Excel attachment...")
    msg = MIMEMultipart()
    msg["From"] = email_from
    msg["To"] = email_from
    n_act = len(results.get('actionable', []))
    n_wl = len(results.get('watchlist', []))
    n_ex = sum(1 for t in results.get('recent', []) if t.get('fresh_exit'))
    n_op = sum(1 for t in results.get('recent', []) if t.get('is_open'))
    n_ultra = sum(1 for s in results.get('actionable', []) if s.get('tier')=='ULTRA')
    msg["Subject"] = f"📈 BMR: {n_ex} EXIT + {n_op} Open + {n_act} New ({n_ultra}💎) + {n_wl} Radar | {now_ist().strftime('%d-%b-%Y')}"

    msg.attach(MIMEText(html_report, "html", "utf-8"))

    with open(excel_file, "rb") as f:
        part = MIMEApplication(f.read(), _subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    part.add_header("Content-Disposition", "attachment", filename=excel_file)
    msg.attach(part)

    # Gmail can transiently reset connections from hosted runners.  Use a
    # bounded retry and an explicit socket timeout so a run cannot hang until
    # GitHub kills it with no useful diagnosis.
    last_error = None
    for attempt in range(1, 4):
        try:
            with smtplib.SMTP("smtp.gmail.com", 587, timeout=45) as server:
                server.ehlo()
                server.starttls()
                server.ehlo()
                server.login(email_from, app_password)
                server.send_message(msg)
            print("✅ Email sent successfully with Excel attachment!")
            break
        except (smtplib.SMTPException, OSError) as exc:
            last_error = exc
            print(f"[!] SMTP attempt {attempt}/3 failed: {exc}")
            if attempt < 3:
                time.sleep(5 * attempt)
    else:
        raise RuntimeError(f"Could not send daily report after 3 SMTP attempts: {last_error}") from last_error
else:
    # Do not report a green workflow when the actual deliverable was never
    # sent.  The artifact step still preserves diagnostics on failed runs.
    raise RuntimeError("MY_EMAIL and MY_APP_PASSWORD must be configured; refusing to silently skip the daily report.")
